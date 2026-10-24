import streamlit as st
import pandas as pd
import numpy as np
import altair as alt
import os, re, json, math, datetime, html, io, unicodedata
from collections import Counter

st.set_page_config(page_title="Aiversecity - AI Career Navigator", layout="wide")

DATA_FILE = "ai_jobs_market_2025_2026.csv"
CACHE_FILE = "trends_cache.json"
KAGGLE_URL = "https://www.kaggle.com/datasets/alitaqishah/ai-jobs-market-2025-2026-salaries"
ARXIV_RSS = "https://rss.arxiv.org/rss/cs.AI"

# [NEW] 实时招聘信号（Arbeitnow 公开 API）
LIVE_JOBS_URL = "https://www.arbeitnow.com/api/job-board-api"
LIVE_CACHE = "live_jobs_cache.json"
LIVE_MOM_W = 0.05   # 实时占比(%)进入技能动量的权重（假设值）
LIVE_W = 0.5        # 实时需求进入岗位排序微调的权重（假设值，只影响排序）
LIVE_MIN = 5.0      # 技能在实时广告中占比 ≥ 此值(%) 才算“有实时需求”
LIVE_SHARE, LIVE_N, LIVE_META, LIVE_CACHED = {}, 0, {}, True

UP, DOWN, FLAT = "UP", "DOWN", "FLAT"  # 语言无关的内部状态码，显示时用 status_label() 转换

# =====================================================================
# 多语言（中文 / English / Bahasa Malaysia）
# 每条文字是 (中文, English, Bahasa Malaysia)
# =====================================================================
LANGS = {"English": "en", "Bahasa Malaysia": "ms", "中文": "zh"}
_IDX = {"zh": 0, "en": 1, "ms": 2}
LANG = "zh"

TXT = {
    "lite": ("📶 低流量模式（不加载图表和实时抓取）", "📶 Low-bandwidth mode (no charts, no live fetch)",
             "📶 Mod jalur lebar rendah (tiada carta, tiada ambilan langsung)"),
    "app_title": ("🚀 Aiversecity - AI 时代职业与技能导航仪", "🚀 Aiversecity – AI Career & Skills Navigator",
                  "🚀 Aiversecity – Pemandu Kerjaya & Kemahiran Era AI"),
    "human_note": (
        "🤝 **以人为本：** 本工具是支持人类职业决策的参考工具，不会替代人的判断，也不会做出任何最终招聘、录用或晋升决定。所有结果仅供规划参考。",
        "🤝 **Human-centred:** This tool supports human career decisions. It does not replace human judgement and never makes final hiring, selection or promotion decisions. All results are for planning reference only.",
        "🤝 **Berpusatkan manusia:** Alat ini menyokong keputusan kerjaya manusia. Ia tidak menggantikan pertimbangan manusia dan tidak membuat keputusan akhir pengambilan, pemilihan atau kenaikan pangkat. Semua hasil adalah untuk rujukan perancangan sahaja."),
    # ---- 侧边栏 ----
    "sb_signals": ("📊 AI 趋势信号层", "📊 AI Trend Signals", "📊 Isyarat Trend AI"),
    "sb_src1": ("① 招聘数据（开放数据集）", "① Job-market data (open dataset)", "① Data pasaran pekerjaan (set data terbuka)"),
    "sb_missing": ("未找到 {f}", "{f} not found", "{f} tidak dijumpai"),
    "sb_source": ("**来源：** [Kaggle · AI Jobs Market 2025-2026]({url})", "**Source:** [Kaggle · AI Jobs Market 2025-2026]({url})",
                  "**Sumber:** [Kaggle · AI Jobs Market 2025-2026]({url})"),
    "sb_filedate": ("**数据文件最后修改日期：** {d}（不是数据集发布日期）",
                    "**Data file last modified:** {d} (not the dataset's publication date)",
                    "**Tarikh fail diubah suai:** {d} (bukan tarikh penerbitan set data)"),
    "sb_period": ("**数据覆盖的招聘时间：** {r}", "**Postings period covered:** {r}",
                  "**Tempoh iklan diliputi:** {r}"),
    "sb_samples": ("**样本数：** {n} 条招聘记录", "**Samples:** {n} job-posting records", "**Sampel:** {n} rekod iklan jawatan"),
    "sb_src2": ("② 最新 AI 研究动态（实时更新）", "② Latest AI research (live)", "② Penyelidikan AI terkini (langsung)"),
    "sb_pub": ("发布：{p}", "Published: {p}", "Diterbitkan: {p}"),
    "sb_arxiv": ("来源：arXiv cs.AI RSS · 抓取时间：{ts}", "Source: arXiv cs.AI RSS · fetched: {ts}", "Sumber: arXiv cs.AI RSS · diambil: {ts}"),
    "cached": (" （缓存）", " (cached)", " (cache)"),
    "sb_nosig": ("实时来源暂不可用，且暂无缓存。", "Live source unavailable and no cache yet.", "Sumber langsung tidak tersedia dan tiada cache."),
    "sb_disc_title": ("🤖 AI 披露与数据说明", "🤖 AI Disclosure & Data", "🤖 Pendedahan AI & Data"),
    "sb_disc": (
        "- **未使用任何大语言模型（LLM）。** 岗位匹配、趋势雷达、SDG 映射、学习计划全部由规则和公开数据计算得出。\n"
        "- 数据：Kaggle 公开数据集、arXiv 公开 RSS、公开学习资源链接。\n"
        "- 不使用真实个人数据；示例成员均为合成数据。\n"
        "- 系统不会自行编造岗位、课程或统计数字；数据缺失时会明确显示“无数据”。",
        "- **No large language model (LLM) is used.** Role matching, the trend radar, SDG mapping and learning plans are all computed from rules and public data.\n"
        "- Data: the public Kaggle dataset, the public arXiv RSS feed and public learning-resource links.\n"
        "- No real personal data is used; sample members are synthetic.\n"
        "- The system never makes up jobs, courses or statistics; when data is missing it says so.",
        "- **Tiada model bahasa besar (LLM) digunakan.** Padanan jawatan, radar trend, pemetaan SDG dan pelan pembelajaran semuanya dikira daripada peraturan dan data awam.\n"
        "- Data: set data awam Kaggle, suapan RSS awam arXiv dan pautan sumber pembelajaran awam.\n"
        "- Tiada data peribadi sebenar digunakan; ahli sampel adalah data sintetik.\n"
        "- Sistem tidak mereka jawatan, kursus atau statistik; apabila data tiada, ia dinyatakan dengan jelas."),
    "sb_assump": ("⚙️ 设置与情景假设（可调）", "⚙️ Settings & scenario assumptions (adjustable)",
                  "⚙️ Tetapan & andaian senario (boleh laras)"),
    "sub_path": ("🧭 我的路径", "🧭 My path", "🧭 Laluan saya"),
    "sub_trend": ("📡 市场趋势", "📡 Market trends", "📡 Trend pasaran"),
    "sub_trust": ("⚖️ 公平与来源", "⚖️ Fairness & sources", "⚖️ Keadilan & sumber"),
    "pick_role": ("选择要查看的岗位（按匹配度排序）", "Choose a role to view (ranked by match)", "Pilih jawatan untuk dilihat (disusun mengikut padanan)"),
    "why_h": ("📎 得分构成、数据来源与需求情景", "📎 Score breakdown, sources & demand scenarios", "📎 Pecahan skor, sumber & senario permintaan"),
    "app_sub": ("AI 时代职业与技能导航仪", "AI Career & Skills Navigator", "Pemandu Kerjaya & Kemahiran Era AI"),
    "hero_pts": ("看清趋势|匹配技能|补齐差距|免费学习", "See the trends|Match your skills|Close the gaps|Learn for free",
                 "Lihat trend|Padankan kemahiran|Tutup jurang|Belajar percuma"),
    "tagline": ("对照你的技能与 AI 岗位的走向，再用免费资源一步步补齐差距。",
                "See how your skills line up with where AI jobs are heading, then close the gaps with free learning resources.",
                "Lihat bagaimana kemahiran anda selari dengan arah pekerjaan AI, kemudian tutup jurang dengan sumber pembelajaran percuma."),
    "mult_slow": ("缓慢采用：增长倍数", "Slow adoption: growth multiplier", "Penerimaan perlahan: pengganda pertumbuhan"),
    "mult_mod": ("温和变革：增长倍数", "Moderate change: growth multiplier", "Perubahan sederhana: pengganda pertumbuhan"),
    "mult_fast": ("快速颠覆：增长倍数", "Rapid disruption: growth multiplier", "Gangguan pesat: pengganda pertumbuhan"),
    "spread": ("不确定区间 ±%", "Uncertainty range ±%", "Julat ketidakpastian ±%"),
    "skills_only": ("⚖️ 仅按技能匹配（不使用背景相关度）", "⚖️ Match on skills only (ignore background)",
                    "⚖️ Padanan berdasarkan kemahiran sahaja (abaikan latar belakang)"),
    "err_nodata": ("缺少数据文件 {f}。请先下载 Kaggle 数据集并放到同目录。",
                   "Missing data file {f}. Download the Kaggle dataset and put it in the same folder.",
                   "Fail data {f} tiada. Muat turun set data Kaggle dan letak dalam folder yang sama."),
    "err_cols": ("CSV 列名无法识别（需要岗位名列和技能列）。当前列名：",
                 "Unrecognised CSV columns (need a job-title column and a skills column). Current columns:",
                 "Lajur CSV tidak dikenali (perlu lajur jawatan dan lajur kemahiran). Lajur semasa:"),
    "tab_me": ("👤 个人模式", "👤 Individual mode", "👤 Mod individu"),
    "tab_team": ("👥 团队 / 批量模式（HR · 职业中心）", "👥 Team / batch mode (HR · Career centre)",
                 "👥 Mod pasukan / kelompok (HR · Pusat kerjaya)"),
    # ---- 个人录入 ----
    "h_intake": ("👤 你的技能画像", "👤 Your skills profile", "👤 Profil kemahiran anda"),
    "in_role": ("你当前的职业/专业背景：", "Your current job / field of study:", "Pekerjaan / bidang pengajian semasa anda:"),
    "in_skills": ("你掌握的技能（逗号分隔）：", "Your skills (comma-separated):", "Kemahiran anda (dipisahkan koma):"),
    "in_cv": ("（可选）上传 CV：PDF / DOCX / TXT，系统会自动识别技能",
              "(Optional) Upload a CV: PDF / DOCX / TXT — skills are detected automatically",
              "(Pilihan) Muat naik CV: PDF / DOCX / TXT — kemahiran dikesan secara automatik"),
    "cv_notice": ("🔒 Demo 请只上传合成/示例 CV，**不要上传真实个人 CV**。文件只在本次会话内存中处理，不会被保存。",
                  "🔒 For this demo, upload only synthetic/sample CVs — **do not upload a real personal CV**. Files are processed in memory for this session only and are not stored.",
                  "🔒 Untuk demo ini, muat naik CV sampel/sintetik sahaja — **jangan muat naik CV peribadi sebenar**. Fail diproses dalam memori sesi ini sahaja dan tidak disimpan."),
    "q_title": ("**潜力问卷**", "**Potential questionnaire**", "**Soal selidik potensi**"),
    "q_hours": ("每周可投入学习的小时数", "Hours per week you can study", "Jam seminggu untuk belajar"),
    "q_pref": ("你更倾向的学习方式", "Preferred learning style", "Gaya pembelajaran pilihan"),
    "pref_self": ("自学为主（视频 + 练习）", "Mostly self-study (videos + practice)", "Kebanyakan belajar sendiri (video + latihan)"),
    "pref_proj": ("项目驱动", "Project-driven", "Berasaskan projek"),
    "pref_course": ("课程 + 证书", "Courses + certificates", "Kursus + sijil"),
    "q_mind": ("你对学习全新技术的意愿", "Willingness to learn new technology", "Kesediaan mempelajari teknologi baharu"),
    "mind_1": ("一般", "Moderate", "Sederhana"),
    "mind_2": ("较高", "High", "Tinggi"),
    "mind_3": ("很高", "Very high", "Sangat tinggi"),
    "q_sdg": ("你最关心哪些 SDG？（可多选，会提高相关岗位的排名）",
              "Which SDGs matter most to you? (multi-select; boosts related roles in the ranking)",
              "SDG mana yang paling penting bagi anda? (pilih beberapa; menaikkan kedudukan jawatan berkaitan)"),
    "sdg_pref_on": ("已应用你的 SDG 偏好：每个重合的 SDG 在排序中加 {w} 分（只影响排序，不改变匹配度分数）。",
                    "Your SDG preference is applied: each overlapping SDG adds {w} points to the ranking only (the match score is unchanged).",
                    "Pilihan SDG anda digunakan: setiap SDG yang bertindih menambah {w} mata pada susunan sahaja (skor padanan tidak berubah)."),
    "sdg_match": ("此岗位符合你关心的：{s}", "This role matches your chosen SDGs: {s}",
                  "Jawatan ini sepadan dengan SDG pilihan anda: {s}"),
    "sdg_bonus_line": ("SDG 偏好加权（仅用于排序）：{b:+.0f}", "SDG preference weighting (ranking only): {b:+.0f}",
                       "Pemberat pilihan SDG (susunan sahaja): {b:+.0f}"),
    "cv_found": ("从 CV 识别到 {n} 项技能：{s}", "Detected {n} skills from the CV: {s}", "{n} kemahiran dikesan daripada CV: {s}"),
    "cv_none": ("CV 已读取，但没有识别到数据集中的技能关键词。", "CV read, but no skills from the dataset were detected.",
                "CV dibaca, tetapi tiada kemahiran daripada set data dikesan."),
    "cv_fail": ("CV 读取失败：{e}", "Could not read the CV: {e}", "Gagal membaca CV: {e}"),
    # ---- 雷达 ----
    "h_radar": ("📡 未来技能雷达", "📡 Future Skills Radar", "📡 Radar Kemahiran Masa Depan"),
    "radar_nodata": ("数据集缺少 posting_year 列，无法计算上升 / 下降趋势。",
                     "The dataset has no posting_year column, so rising/cooling trends cannot be computed.",
                     "Set data tiada lajur posting_year, trend meningkat/menurun tidak dapat dikira."),
    "radar_method": (
        "方法：比较同一数据集中 2025 年（{n25} 条）与 2026 年（{n26} 条）招聘记录里，各岗位 / 各技能所占份额的变化。蓝色 = 份额上升，橙色 = 份额下降（相对下降，不等于招聘总量减少）。注意：2026 年数据只覆盖 1–3 月，且单个岗位样本量较小（每年约 20–50 条），结论仅为趋势信号，不是预测。",
        "Method: compare the share of each role / skill among the {n25} job records of 2025 and the {n26} records of 2026 in the same dataset. Blue = share rising, orange = share falling (relative cooling, not necessarily fewer jobs). Note: 2026 data only covers Jan–Mar and each role has a small sample (about 20–50 records per year), so results are trend signals, not forecasts.",
        "Kaedah: bandingkan bahagian setiap jawatan / kemahiran antara {n25} rekod 2025 dan {n26} rekod 2026 dalam set data yang sama. Biru = bahagian meningkat, jingga = bahagian menurun (penyejukan relatif, bukan semestinya pekerjaan berkurang). Nota: data 2026 hanya meliputi Jan–Mac dan sampel setiap jawatan kecil (kira-kira 20–50 rekod setahun), jadi keputusan ialah isyarat trend, bukan ramalan."),
    "radar_roles": ("岗位：份额上升 / 下降", "Roles: rising / cooling share", "Jawatan: bahagian meningkat / menurun"),
    "radar_skills": ("技能：份额上升 / 下降（各取前 10）", "Skills: rising / cooling share (top 10 each)",
                     "Kemahiran: bahagian meningkat / menurun (10 teratas setiap satu)"),
    "axis_shift": ("招聘占比变化（百分点，2025 → 2026）", "Change in share of postings (percentage points, 2025 → 2026)",
                   "Perubahan bahagian iklan (mata peratusan, 2025 → 2026)"),
    "radar_exp": ("📋 查看完整数据表（岗位 / 技能）", "📋 Full data tables (roles / skills)", "📋 Jadual data penuh (jawatan / kemahiran)"),
    "col_role": ("岗位", "Role", "Jawatan"),
    "col_skill": ("技能", "Skill", "Kemahiran"),
    "col_trend": ("趋势", "Trend", "Trend"),
    "col_shift": ("份额变化(pp)", "Share change (pp)", "Perubahan bahagian (pp)"),
    "col_n25r": ("2025 记录数", "2025 records", "Rekod 2025"),
    "col_n26r": ("2026 记录数", "2026 records", "Rekod 2026"),
    "col_n25s": ("2025 提及数", "2025 mentions", "Sebutan 2025"),
    "col_n26s": ("2026 提及数", "2026 mentions", "Sebutan 2026"),
    "col_g_role": ("数据集年增长率(%)", "Dataset YoY growth (%)", "Pertumbuhan YoY set data (%)"),
    "col_g_skill": ("相关岗位平均增长率(%)", "Avg growth of related roles (%)", "Purata pertumbuhan jawatan berkaitan (%)"),
    "up": ("📈 上升", "📈 Rising", "📈 Meningkat"),
    "down": ("📉 下降", "📉 Cooling", "📉 Menurun"),
    "flat": ("➖ 持平", "➖ Flat", "➖ Mendatar"),
    "your_pos": ("**🧭 你的技能在雷达上的位置**", "**🧭 Where your skills sit on the radar**", "**🧭 Kedudukan kemahiran anda pada radar**"),
    "pos_up": ("📈 上升：{s}", "📈 Rising: {s}", "📈 Meningkat: {s}"),
    "pos_down": ("📉 下降：{s}", "📉 Cooling: {s}", "📉 Menurun: {s}"),
    "pos_flat": ("➖ 持平：{s}", "➖ Flat: {s}", "➖ Mendatar: {s}"),
    "none": ("暂无", "none", "tiada"),
    "col_heat": ("arXiv 研究热度（篇数/总篇数）", "arXiv research heat (papers/total)", "Haba penyelidikan arXiv (kertas/jumlah)"),
    "col_mom": ("动量分", "Momentum score", "Skor momentum"),
    "radar_mom_title": ("🔥 技能动量榜（招聘份额变化 + arXiv 研究热度）", "🔥 Skill momentum ranking (job-share change + arXiv research heat)",
                        "🔥 Ranking momentum kemahiran (perubahan bahagian pekerjaan + haba penyelidikan arXiv)"),
    "radar_mom_cap": ("动量分 = 招聘份额变化(pp) + {w} × 研究热度（占最新论文的 %）。权重 {w} 是我们设定的假设值；研究热度反映学术关注度，不等于招聘需求。",
                      "Momentum = job-share change (pp) + {w} × research heat (% of the latest papers). The weight {w} is our assumption; research heat reflects academic interest, not hiring demand.",
                      "Momentum = perubahan bahagian pekerjaan (pp) + {w} × haba penyelidikan (% kertas terkini). Pemberat {w} ialah andaian kami; haba penyelidikan mencerminkan minat akademik, bukan permintaan pekerjaan."),
    "plan_why": ("排序依据（招聘份额变化 ｜ 提到该技能的论文数/总数）：{s}",
                 "Ordering basis (job-share change | papers mentioning the skill / total): {s}",
                 "Asas susunan (perubahan bahagian pekerjaan | kertas menyebut kemahiran / jumlah): {s}"),
    "pos_unknown": (
        "以下技能不在数据集的技能词表中（或样本太少），未给出趋势：{s}。这不代表它们没有价值，只是该数据集没有覆盖。",
        "These skills are not in the dataset's skill vocabulary (or have too few samples), so no trend is shown: {s}. This does not mean they lack value — the dataset simply does not cover them.",
        "Kemahiran berikut tiada dalam senarai kemahiran set data (atau sampel terlalu sedikit), jadi trend tidak ditunjukkan: {s}. Ini tidak bermakna ia tidak bernilai — set data tidak meliputinya."),
    # ---- 匹配 ----
    "btn_run": ("🔥 生成我的未来职业匹配与再培训路径", "🔥 Generate my future-career matches & reskilling path",
                "🔥 Jana padanan kerjaya masa depan & laluan latihan semula"),
    "h_match": ("🎯 你的岗位匹配", "🎯 Your role matches", "🎯 Padanan jawatan anda"),
    "match_cap": (
        "匹配度 = {cw}% × 技能覆盖率 + {bw}% × 背景相关度。趋势信号只用于排序微调，不改变匹配度分数。所有预测均为估计，不构成就业保证。",
        "Match score = {cw}% × skill coverage + {bw}% × background relevance. Trend signals only fine-tune the ranking and never change the score. All projections are estimates, not employment guarantees.",
        "Skor padanan = {cw}% × liputan kemahiran + {bw}% × kaitan latar belakang. Isyarat trend hanya melaraskan susunan dan tidak mengubah skor. Semua unjuran adalah anggaran, bukan jaminan pekerjaan."),
    "replan_changed": (
        "🔄 **自动重新规划：** 加入最新趋势信号（雷达 + arXiv）后，推荐发生了变化。原排序：{old} → 现排序：{new}",
        "🔄 **Automatic re-planning:** after adding the latest trend signals (radar + arXiv), the recommendations changed. Before: {old} → Now: {new}",
        "🔄 **Perancangan semula automatik:** selepas menambah isyarat trend terkini (radar + arXiv), cadangan berubah. Sebelum: {old} → Sekarang: {new}"),
    "replan_same": (
        "🔄 **自动重新规划：** 已应用最新趋势信号，当前推荐顺序没有变化。学习计划会优先安排雷达上“上升”的技能。",
        "🔄 **Automatic re-planning:** the latest trend signals were applied and the recommended order did not change. Learning plans prioritise skills that are rising on the radar.",
        "🔄 **Perancangan semula automatik:** isyarat trend terkini telah digunakan dan susunan cadangan tidak berubah. Pelan pembelajaran mengutamakan kemahiran yang meningkat pada radar."),
    "rec_role": ("🎯 推荐岗位：{t}", "🎯 Recommended role: {t}", "🎯 Jawatan disyorkan: {t}"),
    "m_score": ("匹配度得分", "Match score", "Skor padanan"),
    "m_postings": ("招聘记录数", "Job records", "Rekod pekerjaan"),
    "m_demand": ("市场需求指数", "Demand index", "Indeks permintaan"),
    "m_trend": ("雷达趋势", "Radar trend", "Trend radar"),
    "trend_help": ("2025→2026 招聘占比变化 {v:+.2f}pp", "Change in share of postings 2025→2026: {v:+.2f}pp",
                   "Perubahan bahagian iklan 2025→2026: {v:+.2f}pp"),
    "yes": ("有", "yes", "ada"),
    "no": ("无", "no", "tiada"),
    "score_break": ("得分构成：技能覆盖率 {c:.0f}% × {cw}% + 背景相关度 {b} × {bw}%　｜　趋势加权（仅用于排序）：{tb:+.1f}",
                    "Score breakdown: skill coverage {c:.0f}% × {cw}% + background relevance {b} × {bw}%　|　trend weighting (ranking only): {tb:+.1f}",
                    "Pecahan skor: liputan kemahiran {c:.0f}% × {cw}% + kaitan latar belakang {b} × {bw}%　|　pemberat trend (susunan sahaja): {tb:+.1f}"),
    "src_job": ("📚 来源：Kaggle AI Jobs Market 2025-2026 开放数据集（该岗位 {n} 条招聘记录）。",
                "📚 Source: Kaggle AI Jobs Market 2025-2026 open dataset ({n} job records for this role).",
                "📚 Sumber: set data terbuka Kaggle AI Jobs Market 2025-2026 ({n} rekod untuk jawatan ini)."),
    "sig_hit": ("🔥 趋势信号：最近 {n} 篇 arXiv cs.AI 论文中有 {h} 篇与该岗位关键词相关（抓取于 {ts}）。这是研究热度信号，不代表招聘需求。",
                "🔥 Trend signal: {h} of the latest {n} arXiv cs.AI papers relate to this role's keywords (fetched {ts}). This is a research-interest signal, not hiring demand.",
                "🔥 Isyarat trend: {h} daripada {n} kertas arXiv cs.AI terkini berkaitan kata kunci jawatan ini (diambil {ts}). Ini isyarat minat penyelidikan, bukan permintaan pengambilan."),
    "sig_none": ("趋势信号：最近 {n} 篇 arXiv 论文中暂无明显相关动态",
                 "Trend signal: none of the latest {n} arXiv papers clearly relate to this role",
                 "Isyarat trend: tiada kertas arXiv terkini ({n}) yang jelas berkaitan jawatan ini"),
    "have_skills": ("**已具备的相关技能：** {s}", "**Relevant skills you already have:** {s}", "**Kemahiran berkaitan yang anda ada:** {s}"),
    "generic": ("通用基础能力", "general foundation skills", "kemahiran asas am"),
    "gap_skills": ("**技能差距 (Skills Gap)：** {s}", "**Skills Gap:** {s}", "**Jurang Kemahiran:** {s}"),
    "no_gap": ("暂无明显差距，建议提升实践经验", "No obvious gap — build more hands-on experience",
               "Tiada jurang ketara — tingkatkan pengalaman praktikal"),
    "scen_title": ("**📈 未来需求情景（假设情景，基准：数据集增长率 {g:+.1f}%）**",
                   "**📈 Future demand scenarios (assumption-based; baseline: dataset growth {g:+.1f}%)**",
                   "**📈 Senario permintaan masa depan (berasaskan andaian; asas: pertumbuhan set data {g:+.1f}%)**"),
    "scen_slow": ("缓慢采用 Slow", "Slow adoption", "Penerimaan perlahan"),
    "scen_mod": ("温和变革 Moderate", "Moderate change", "Perubahan sederhana"),
    "scen_fast": ("快速颠覆 Fast", "Rapid disruption", "Gangguan pesat"),
    "scen_help": ("假设值（数据集增长率 × 可调倍数）", "Assumption (dataset growth × adjustable multiplier)",
                  "Andaian (pertumbuhan set data × pengganda boleh laras)"),
    "scen_range": ("假设区间：{lo:+.1f}% ~ {hi:+.1f}%", "Assumed range: {lo:+.1f}% ~ {hi:+.1f}%", "Julat andaian: {lo:+.1f}% ~ {hi:+.1f}%"),
    "scen_warn": ("⚠️ 倍数和区间是可调的假设，不是统计结果；真实结果取决于技术进展、政策与经济环境。",
                  "⚠️ The multipliers and ranges are adjustable assumptions, not statistics; real outcomes depend on technology, policy and the economy.",
                  "⚠️ Pengganda dan julat ialah andaian boleh laras, bukan statistik; hasil sebenar bergantung pada teknologi, dasar dan ekonomi."),
    "no_growth": ("数据集中没有该岗位的增长率数据，因此不提供需求预测（系统不会编造数字）。",
                  "The dataset has no growth figure for this role, so no demand projection is shown (the system never makes up numbers).",
                  "Set data tiada angka pertumbuhan untuk jawatan ini, maka tiada unjuran permintaan ditunjukkan (sistem tidak mereka angka)."),
    "sdg_chain": ("🌱 **SDG 贡献链：** 你的技能 → **{title}** → **{sdg}**\n\n- {desc}",
                  "🌱 **SDG contribution chain:** your skills → **{title}** → **{sdg}**\n\n- {desc}",
                  "🌱 **Rantaian sumbangan SDG:** kemahiran anda → **{title}** → **{sdg}**\n\n- {desc}"),
    "sdg_chain_team": ("🌱 **SDG 贡献链：** 团队技能 → **{title}** → **{sdg}**\n\n- {desc}",
                       "🌱 **SDG contribution chain:** team skills → **{title}** → **{sdg}**\n\n- {desc}",
                       "🌱 **Rantaian sumbangan SDG:** kemahiran pasukan → **{title}** → **{sdg}**\n\n- {desc}"),
    "sdg_industry": ("\n- 该岗位在数据集中最常见的行业：{s}", "\n- Most common industries for this role in the dataset: {s}",
                     "\n- Industri paling biasa untuk jawatan ini dalam set data: {s}"),
    # ---- 再培训路径 + 激励 ----
    "plan_title": ("📖 针对【{title}】的个性化再培训路径（约 {total} 小时 ≈ {weeks} 周）",
                   "📖 Personalised reskilling path for [{title}] (≈ {total} hours ≈ {weeks} weeks)",
                   "📖 Laluan latihan semula peribadi untuk [{title}] (≈ {total} jam ≈ {weeks} minggu)"),
    "plan_time": ("⏱️ **估计耗时：** 约 {total} 小时；按每周 {wh} 小时约 **{weeks} 周**（假设值，因人而异）",
                  "⏱️ **Estimated time:** about {total} hours; at {wh} hours/week ≈ **{weeks} weeks** (assumption; varies by person)",
                  "⏱️ **Anggaran masa:** kira-kira {total} jam; pada {wh} jam/minggu ≈ **{weeks} minggu** (andaian; berbeza mengikut individu)"),
    "plan_prior": ("已按雷达趋势排序：优先补“上升”中的技能。", "Ordered by radar trend: skills that are rising come first.",
                   "Disusun mengikut trend radar: kemahiran yang meningkat didahulukan."),
    "plan_skill": ("补齐「{sk}」：[{name}]({url}) · {cost} · 约 {h} 小时　〔SDG {sdg}〕",
                   "Learn “{sk}”: [{name}]({url}) · {cost} · about {h} hours　〔SDG {sdg}〕",
                   "Pelajari “{sk}”: [{name}]({url}) · {cost} · kira-kira {h} jam　〔SDG {sdg}〕"),
    "plan_nores": ("以下技能暂无匹配的免费资源，系统不会编造课程：{s}。建议自行搜索或请教导师。",
                   "No matching free resource for these skills; the system will not invent a course: {s}. Please search on your own or ask a mentor.",
                   "Tiada sumber percuma yang sepadan untuk kemahiran ini; sistem tidak akan mereka kursus: {s}. Sila cari sendiri atau tanya mentor."),
    "plan_empty": ("该岗位的缺口技能都没有匹配的资源，下面只保留项目和简历步骤。",
                   "None of this role's gap skills has a matching resource, so only the project and CV steps remain.",
                   "Tiada kemahiran jurang untuk jawatan ini mempunyai sumber sepadan, jadi hanya langkah projek dan CV kekal."),
    "plan_project": ("实操项目：用 Kaggle 公开数据集完成 1 个与「{title}」相关的小项目并放到 GitHub · 约 {h} 小时",
                     "Hands-on project: build one small project related to “{title}” with a public Kaggle dataset and publish it on GitHub · about {h} hours",
                     "Projek praktikal: bina satu projek kecil berkaitan “{title}” dengan set data awam Kaggle dan terbitkan di GitHub · kira-kira {h} jam"),
    "plan_cv": ("简历重构与投递：把新技能和项目写进 CV · 约 {h} 小时", "Rework your CV: add the new skills and project · about {h} hours",
                "Susun semula CV: tambah kemahiran dan projek baharu · kira-kira {h} jam"),
    "plan_cap": ("学习方式偏好：{p}；学习意愿：{g}（仅作参考）", "Learning preference: {p}; willingness: {g} (for reference only)",
                 "Pilihan pembelajaran: {p}; kesediaan: {g} (rujukan sahaja)"),
    "plan_note": ("课程时长、项目和简历耗时均为粗略假设值，请以各课程官方页面标示的时长为准。",
                  "Course, project and CV hours are rough assumptions — check each course's official page for its real duration.",
                  "Jam kursus, projek dan CV ialah andaian kasar — sila semak halaman rasmi setiap kursus untuk tempoh sebenar."),
    "c_free": ("免费", "Free", "Percuma"),
    "c_mostly": ("多数免费", "Mostly free", "Kebanyakannya percuma"),
    "c_audit": ("免费旁听/付费证书", "Free to audit / paid certificate", "Percuma untuk audit / sijil berbayar"),
    "prog_title": ("**🏆 学习进度与徽章**", "**🏆 Progress & badges**", "**🏆 Kemajuan & lencana**"),
    "prog_text": ("进度：{d}/{n} 步完成 · 剩余约 {h} 小时", "Progress: {d}/{n} steps done · about {h} hours left",
                  "Kemajuan: {d}/{n} langkah selesai · baki kira-kira {h} jam"),
    "badge_skill": ("🏷️ 技能徽章：{s}", "🏷️ Skill badge: {s}", "🏷️ Lencana kemahiran: {s}"),
    "badge_start": ("🌱 起步者（完成第 1 步）", "🌱 Starter (first step done)", "🌱 Permulaan (langkah pertama selesai)"),
    "badge_builder": ("🚀 技能建设者（学习步骤全部完成）", "🚀 Skill Builder (all learning steps done)",
                      "🚀 Pembina Kemahiran (semua langkah pembelajaran selesai)"),
    "badge_ready": ("🏅 求职就绪（全部步骤完成）", "🏅 Career Ready (all steps done)", "🏅 Sedia Kerjaya (semua langkah selesai)"),
    "badge_none": ("完成第一步即可获得徽章。", "Finish your first step to earn a badge.", "Selesaikan langkah pertama untuk mendapat lencana."),
    "sdg_overall": ("🌍 **你的职业路径整体贡献的 SDG：** {s}", "🌍 **SDGs your career path contributes to overall:** {s}",
                    "🌍 **SDG yang disumbangkan oleh laluan kerjaya anda:** {s}"),
    "hrd": (
        "HRD Corp e-LATiH：HRD Corp 的在线学习平台，面向个人开放，官方 FAQ 提到提供大量免费课程，另有付费课程：[FAQ](https://supportcentre.hrdcorp.gov.my/portal/en/kb/articles/faqs-hrd-corp-e-latih)。若你已在职，可请雇主查询 [HRD Corp Claimable Courses](https://hrdcorp.gov.my/hrdcorp-claimable-courses)（由雇主申请，需符合其资格条件）。课程数量与资格以官方页面为准。",
        "HRD Corp e-LATiH: HRD Corp's online learning platform, open to individuals; its official FAQ mentions many free courses plus paid ones: [FAQ](https://supportcentre.hrdcorp.gov.my/portal/en/kb/articles/faqs-hrd-corp-e-latih). If you are employed, ask your employer about [HRD Corp Claimable Courses](https://hrdcorp.gov.my/hrdcorp-claimable-courses) (claimed by the employer, subject to eligibility). Course numbers and eligibility follow the official pages.",
        "HRD Corp e-LATiH: platform pembelajaran dalam talian HRD Corp, terbuka kepada individu; FAQ rasminya menyebut banyak kursus percuma serta kursus berbayar: [FAQ](https://supportcentre.hrdcorp.gov.my/portal/en/kb/articles/faqs-hrd-corp-e-latih). Jika anda bekerja, tanya majikan tentang [HRD Corp Claimable Courses](https://hrdcorp.gov.my/hrdcorp-claimable-courses) (dituntut oleh majikan, tertakluk kepada kelayakan). Bilangan kursus dan kelayakan mengikut laman rasmi."),
    # ---- 公平性 ----
    "h_fair": ("⚖️ 公平性检查 (Fairness Check)", "⚖️ Fairness Check", "⚖️ Semakan Keadilan"),
    "fair_cap": (
        "系统不收集也不使用年龄、性别、种族等信息。唯一与“背景”相关的因素是职业/专业关键词（最多占 {bw}% 分数）。下面用同样的技能、换不同背景，检查分数差异。",
        "The system does not collect or use age, gender or ethnicity. The only background-related factor is the job/field keyword (worth at most {bw}% of the score). Below, the same skills are tested with different backgrounds to check for score differences.",
        "Sistem tidak mengumpul atau menggunakan umur, jantina atau etnik. Satu-satunya faktor berkaitan latar belakang ialah kata kunci pekerjaan/bidang (paling banyak {bw}% skor). Di bawah, kemahiran yang sama diuji dengan latar belakang berbeza untuk menyemak perbezaan skor."),
    "fair_role": ("测试岗位（仅按技能时的第一推荐）：**{r}**", "Test role (top match on skills only): **{r}**",
                  "Jawatan ujian (padanan teratas berdasarkan kemahiran sahaja): **{r}**"),
    "fc_bg": ("背景", "Background", "Latar belakang"),
    "fc_score": ("该岗位得分", "Score for this role", "Skor untuk jawatan ini"),
    "fc_top": ("该背景下的首选岗位", "Top role for this background", "Jawatan teratas untuk latar belakang ini"),
    "fair_result": ("同样技能下，不同背景导致的最大分差：**{g} 分**（设计上限 {bw} 分）。技能覆盖率仍是分数的主导因素。",
                    "With identical skills, the largest score gap caused by background is **{g} points** (design cap {bw}). Skill coverage remains the dominant factor.",
                    "Dengan kemahiran yang sama, perbezaan skor terbesar akibat latar belakang ialah **{g} mata** (had reka bentuk {bw}). Liputan kemahiran kekal faktor utama."),
    "fair_zero": ("当前为“仅按技能匹配”模式，背景不影响分数，分差为 0。", "Skills-only mode is on: background does not affect scores (gap = 0).",
                  "Mod kemahiran sahaja aktif: latar belakang tidak menjejaskan skor (perbezaan = 0)."),
    "fair_limit": ("局限：这是基于关键词规则的简单检查，不能证明数据集本身没有偏差。",
                   "Limitation: this is a simple keyword-rule check and cannot prove the dataset itself is free of bias.",
                   "Had: ini semakan peraturan kata kunci yang ringkas dan tidak dapat membuktikan set data bebas daripada bias."),
    "fair_need": ("请先输入技能后再运行公平性检查。", "Enter some skills first to run the fairness check.", "Sila masukkan kemahiran dahulu untuk menjalankan semakan keadilan."),
    # ---- 团队 ----
    "h_team": ("👥 团队 / 批量模式", "👥 Team / batch mode", "👥 Mod pasukan / kelompok"),
    "aud_label": ("你的身份", "Your role", "Peranan anda"),
    "aud_hr": ("HR / L&D 经理：规划团队再培训", "HR / L&D manager: plan team reskilling", "Pengurus HR / L&D: rancang latihan semula pasukan"),
    "aud_tvet": ("TVET / 大学职业中心：为学生批量提供建议", "TVET / university career centre: advise students in batch",
                 "TVET / pusat kerjaya universiti: nasihat pelajar secara kelompok"),
    "team_hours": ("成员平均每周可投入学习小时数", "Average weekly study hours per member", "Purata jam belajar seminggu setiap ahli"),
    "team_human": ("🤝 团队结果仅用于**培训规划与课程设计**，不用于筛选、淘汰、录用或晋升个人；任何决定都应由人来做出。",
                   "🤝 Team results are for **training planning and curriculum design** only — not for screening, rejecting, hiring or promoting individuals. Any decision must be made by people.",
                   "🤝 Keputusan pasukan hanya untuk **perancangan latihan dan reka bentuk kurikulum** — bukan untuk menapis, menolak, mengambil atau menaikkan pangkat individu. Sebarang keputusan mesti dibuat oleh manusia."),
    "enter_members": ("**输入成员 / 学生数据**（上传 CSV，或直接在下方表格编辑）", "**Enter member / student data** (upload a CSV or edit the table below)",
                      "**Masukkan data ahli / pelajar** (muat naik CSV atau edit jadual di bawah)"),
    "dl_tpl": ("⬇️ 下载 CSV 模板", "⬇️ Download CSV template", "⬇️ Muat turun templat CSV"),
    "up_csv": ("上传 CSV（列：name, background, skills, region（可选）；技能用 , ; | 分隔）",
               "Upload CSV (columns: name, background, skills, optional region; separate skills with , ; |)",
               "Muat naik CSV (lajur: name, background, skills, region pilihan; asingkan kemahiran dengan , ; |)"),
    "csv_needs": ("CSV 里需要有 skills（技能）列。", "The CSV needs a skills column.", "CSV memerlukan lajur skills (kemahiran)."),
    "csv_fail": ("CSV 读取失败：{e}", "Could not read the CSV: {e}", "Gagal membaca CSV: {e}"),
    "member": ("成员{i}", "Member {i}", "Ahli {i}"),
    "mem_cap": ("数据只在本次会话内存中处理，不会被保存。请只使用合成数据。", "Data is processed in memory for this session only and is not stored. Use synthetic data only.",
                "Data diproses dalam memori sesi ini sahaja dan tidak disimpan. Gunakan data sintetik sahaja."),
    "need_member": ("请先输入至少一位成员的技能。", "Enter at least one member's skills first.", "Sila masukkan kemahiran sekurang-kurangnya seorang ahli dahulu."),
    "target_sel": ("团队要朝哪个岗位转型？（默认：团队成员最常见的首选岗位）",
                   "Which role should the team transition to? (default: most common top match among members)",
                   "Jawatan mana yang menjadi sasaran peralihan pasukan? (lalai: padanan teratas paling biasa dalam kalangan ahli)"),
    "m_team_avg": ("团队平均匹配度", "Team average match", "Purata padanan pasukan"),
    "m_team_hours": ("团队总学习时长（估计）", "Total team learning hours (estimate)", "Jumlah jam pembelajaran pasukan (anggaran)"),
    "hours_unit": ("{h} 小时", "{h} hours", "{h} jam"),
    "m_target_sig": ("目标岗位市场信号", "Target-role market signal", "Isyarat pasaran jawatan sasaran"),
    "sig_help": ("来自上方雷达：2025→2026 招聘占比变化", "From the radar: change in share of postings 2025→2026",
                 "Daripada radar: perubahan bahagian iklan 2025→2026"),
    "team_cov_title": ("团队技能覆盖：{t}", "Team skill coverage: {t}", "Liputan kemahiran pasukan: {t}"),
    "team_cov_axis": ("团队覆盖率(%)", "Team coverage (%)", "Liputan pasukan (%)"),
    "heat_title": ("🔥 团队技能缺口热力图（成员 × 技能）", "🔥 Team skills-gap heatmap (member × skill)",
                   "🔥 Peta haba jurang kemahiran pasukan (ahli × kemahiran)"),
    "have": ("已具备", "Have", "Ada"),
    "gap": ("缺口", "Gap", "Jurang"),
    "col_member": ("成员", "Member", "Ahli"),
    "prio_title": ("🎯 培训优先级（缺口最多的技能）", "🎯 Training priorities (most common gaps)", "🎯 Keutamaan latihan (jurang paling banyak)"),
    "n_missing": ("缺该技能人数", "Members missing it", "Ahli yang tiada"),
    "pct": ("占比(%)", "Share (%)", "Peratus (%)"),
    "arrange": ("建议安排", "Suggested action", "Tindakan dicadangkan"),
    "group": ("集中培训", "Group training", "Latihan berkumpulan"),
    "indiv": ("个别辅导", "Individual coaching", "Bimbingan individu"),
    "res": ("推荐资源", "Recommended resource", "Sumber disyorkan"),
    "cost": ("费用", "Cost", "Kos"),
    "no_res": ("（暂无匹配资源）", "(no matching resource)", "(tiada sumber sepadan)"),
    "member_detail": ("👤 成员明细", "👤 Member details", "👤 Butiran ahli"),
    "col_bg": ("背景", "Background", "Latar belakang"),
    "match_pct": ("对目标岗位匹配度(%)", "Match to target role (%)", "Padanan kepada jawatan sasaran (%)"),
    "n_gap": ("缺口技能数", "Number of gap skills", "Bilangan kemahiran jurang"),
    "est_h": ("估计学习时长(小时)", "Estimated learning hours", "Anggaran jam belajar"),
    "est_wk": ("估计周数", "Estimated weeks", "Anggaran minggu"),
    "team_cap": ("学习时长为按资源估计值之和（每人最多补 3 个技能 + 项目 8 小时 + 简历 3 小时），均为假设值，仅供规划参考。",
                 "Learning hours are sums of per-resource estimates (up to 3 skills per person + 8 h project + 3 h CV). All are assumptions for planning only.",
                 "Jam pembelajaran ialah jumlah anggaran setiap sumber (sehingga 3 kemahiran setiap orang + 8 j projek + 3 j CV). Semuanya andaian untuk perancangan sahaja."),
    "student": ("学生", "Student", "Pelajar"),
    "top_role": ("首选岗位", "Top role", "Jawatan teratas"),
    "match_col": ("匹配度(%)", "Match (%)", "Padanan (%)"),
    "alts": ("备选岗位", "Alternatives", "Alternatif"),
    "top_gap": ("最大缺口技能", "Biggest gap skills", "Jurang kemahiran terbesar"),
    "n_students": ("学生人数", "Students", "Bilangan pelajar"),
    "avg_match": ("平均首选匹配度", "Average top match", "Purata padanan teratas"),
    "top_dist": ("首选岗位分布", "Top-role distribution", "Taburan jawatan teratas"),
    "count": ("人数", "Count", "Bilangan"),
    "common_gaps": ("全体学生最常见的技能缺口", "Most common skill gaps across students", "Jurang kemahiran paling biasa merentas pelajar"),
    "gap_n": ("缺口人数", "Students with gap", "Pelajar dengan jurang"),
    "per_student": ("👤 逐个学生建议", "👤 Advice per student", "👤 Nasihat setiap pelajar"),
    "dl_result": ("⬇️ 下载结果 CSV", "⬇️ Download results CSV", "⬇️ Muat turun keputusan CSV"),
    "cohort_sdg": ("🌍 **该群体的职业路径整体贡献的 SDG：** {s}", "🌍 **SDGs this cohort's career paths contribute to:** {s}",
                   "🌍 **SDG yang disumbangkan oleh laluan kerjaya kohort ini:** {s}"),
    "curriculum": ("建议课程设计：优先针对上方「最常见的技能缺口」开设短课程（可对应免费资源，如 Kaggle Learn、Google MLCC）。",
                   "Suggested curriculum design: run short courses first on the most common skill gaps above (free resources such as Kaggle Learn or Google MLCC can be used).",
                   "Cadangan reka bentuk kurikulum: adakan kursus pendek dahulu untuk jurang kemahiran paling biasa di atas (boleh guna sumber percuma seperti Kaggle Learn atau Google MLCC)."),
    # ---- 地区 SDG ----
    "h_region": ("🗺️ 地区 SDG 贡献视图", "🗺️ SDG contribution by region", "🗺️ Sumbangan SDG mengikut wilayah"),
    "region": ("地区", "Region", "Wilayah"),
    "n_people": ("人数", "People", "Bilangan"),
    "top_sdg": ("主要贡献的 SDG（人数）", "Top SDGs (people)", "SDG utama (bilangan)"),
    "no_region": ("未填写", "Not specified", "Tidak dinyatakan"),
    "region_hint": ("在 CSV 或表格中添加 region（地区）列，即可看到按地区汇总的 SDG 贡献。",
                    "Add a region column to the CSV or table to see SDG contributions aggregated by region.",
                    "Tambah lajur region (wilayah) pada CSV atau jadual untuk melihat sumbangan SDG mengikut wilayah."),
    "region_cap": ("按每位成员的目标/首选岗位对应的 SDG 汇总，是概念关联，不是官方统计。",
                   "Aggregated from the SDGs linked to each member's target / top role. This is a conceptual link, not an official statistic.",
                   "Dikumpul daripada SDG yang dikaitkan dengan jawatan sasaran/teratas setiap ahli. Ini kaitan konsep, bukan statistik rasmi."),
    # ---- SDG 覆盖 ----
    "h_cov": ("🌍 17 项 SDG 覆盖总览（岗位 / 技能 / 行业应用）", "🌍 Overview of all 17 SDGs (roles / skills / industry applications)",
              "🌍 Gambaran keseluruhan 17 SDG (jawatan / kemahiran / aplikasi industri)"),
    "c_roles": ("关联岗位数", "Linked roles", "Jawatan berkaitan"),
    "c_skills": ("关联技能数", "Linked skills", "Kemahiran berkaitan"),
    "c_ind": ("行业应用层", "Industry layer", "Lapisan industri"),
    "c_cov": ("覆盖情况", "Coverage", "Liputan"),
    "cov_direct": ("直接关联", "Direct link", "Pautan langsung"),
    "cov_ind": ("仅行业应用层", "Industry layer only", "Lapisan industri sahaja"),
    "cov_none": ("本数据集未覆盖", "Not covered by this dataset", "Tidak diliputi set data ini"),
    "cov_cap": ("映射为基于关键词的规则（可在代码 ROLE_RULES / SKILL_RULES 中调整），是合理的概念关联，不是官方分类。海洋、陆地生态等 SDG 与本数据集中的 AI 岗位关联较弱，如实标注为未覆盖，而不是硬凑。",
                "The mapping is keyword-rule based (adjust ROLE_RULES / SKILL_RULES in the code) — a reasonable conceptual link, not an official classification. SDGs such as life below water or on land have weak links to the AI roles in this dataset, so they are honestly marked as not covered instead of forced.",
                "Pemetaan berasaskan peraturan kata kunci (laraskan ROLE_RULES / SKILL_RULES dalam kod) — kaitan konsep yang munasabah, bukan klasifikasi rasmi. SDG seperti kehidupan bawah air atau di darat mempunyai kaitan lemah dengan jawatan AI dalam set data ini, jadi ditandakan tidak diliputi dan tidak dipaksa."),
}

# [NEW] 实时招聘信号的三语文字
TXT.update({
    "live_title": ("实时招聘信号（自动更新，无需用户操作）", "Live job-market signal (auto-updated, no user action)",
                   "Isyarat pasaran pekerjaan langsung (dikemas kini automatik)"),
    "live_src": ("**来源：** [Arbeitnow 公开招聘 API](https://www.arbeitnow.com/)（免费，无需密钥）",
                 "**Source:** [Arbeitnow public job-board API](https://www.arbeitnow.com/) (free, no key)",
                 "**Sumber:** [API papan pekerjaan awam Arbeitnow](https://www.arbeitnow.com/) (percuma, tiada kunci)"),
    "live_cap": ("统计最新招聘广告中提到各技能的比例，并与数据集 2026 年的比例对比。该比例以权重 {w} 加入技能动量（影响学习计划顺序），并对岗位排序做微调（不改变匹配度分数）。",
                 "Share of the latest job ads that mention each skill, compared with the dataset's 2026 share. It is added to skill momentum with weight {w} (affecting learning-plan order) and fine-tunes role ranking (the match score is unchanged).",
                 "Peratus iklan pekerjaan terkini yang menyebut setiap kemahiran, dibandingkan dengan bahagian set data 2026. Ia ditambah kepada momentum kemahiran dengan pemberat {w} (mempengaruhi susunan pelan pembelajaran) dan melaraskan susunan jawatan (skor padanan tidak berubah)."),
    "live_n": ("分析的招聘广告数：{n}", "Job ads analysed: {n}", "Iklan pekerjaan dianalisis: {n}"),
    "live_when": ("抓取时间：{ts}", "Fetched: {ts}", "Diambil: {ts}"),
    "live_range": ("广告发布日期范围：{r}", "Ad posting dates: {r}", "Tarikh iklan: {r}"),
    "live_none": ("实时招聘来源暂不可用，且暂无缓存；系统改用数据集趋势。",
                  "Live job source unavailable and no cache yet; the system falls back to dataset trends.",
                  "Sumber pekerjaan langsung tidak tersedia dan tiada cache; sistem menggunakan trend set data."),
    "live_col_live": ("实时占比(%)", "Live share (%)", "Peratus langsung (%)"),
    "live_col_ds": ("数据集2026占比(%)", "Dataset 2026 share (%)", "Peratus set data 2026 (%)"),
    "live_col_diff": ("差值(pp)", "Difference (pp)", "Perbezaan (pp)"),
    "live_bias": ("局限：该来源以欧洲/德语区（含远程）科技岗位为主，不代表马来西亚市场；只作为“实时需求信号”，不是预测。",
                  "Limitation: this source is mostly European/German-speaking tech jobs (including remote), not the Malaysian market; it is a live demand signal, not a forecast.",
                  "Had: sumber ini kebanyakannya pekerjaan teknologi Eropah/negara berbahasa Jerman (termasuk jarak jauh), bukan pasaran Malaysia; ia isyarat permintaan langsung, bukan ramalan."),
})


# [NEW] CV 识别反馈文字（cv_none 在此覆盖为更详细的说明）
TXT.update({
    "cv_empty": ("没有从 CV 中读到任何文字。如果是扫描版 PDF（图片），系统无法识别，请改用可复制文字的 PDF、DOCX 或 TXT。",
                 "No text could be read from the CV. If it is a scanned (image) PDF the system cannot read it — please use a PDF with selectable text, a DOCX or a TXT file.",
                 "Tiada teks dapat dibaca daripada CV. Jika ia PDF imbasan (imej), sistem tidak dapat membacanya — sila guna PDF dengan teks boleh dipilih, DOCX atau TXT."),
    "cv_chars": ("已读取 CV 文字 {n} 个字符。", "Read {n} characters of text from the CV.", "Membaca {n} aksara teks daripada CV."),
    "cv_none": ("CV 已读取，但没有识别到数据集中的技能词。系统只识别数据集中出现过的技能名称（例如 python、sql、excel），请在左边的技能框里手动补充。",
                "The CV was read, but no skill names from the dataset were found. The system only recognises skill names that appear in the dataset (e.g. python, sql, excel) — please add your skills manually in the skills box.",
                "CV dibaca, tetapi tiada nama kemahiran daripada set data dijumpai. Sistem hanya mengenali nama kemahiran dalam set data (cth. python, sql, excel) — sila tambah kemahiran anda secara manual dalam kotak kemahiran."),
})


# 极简风：去掉标题里的 emoji；趋势用 ▲ ▼ 表示
_ARROWS = {"📈 ": "▲ ", "📉 ": "▼ ", "➖ ": "— "}
_EMOJI_LEAD = re.compile(r"^(\*\*)?(?:[\U0001F000-\U0001FFFF\u2600-\u27BF\uFE0F\u200d]+\s*)+")


def _plain(v):
    for k, a in _ARROWS.items():
        v = v.replace(k, a)
    return _EMOJI_LEAD.sub(lambda m: m.group(1) or "", v)


TXT = {k: tuple(_plain(x) for x in v) for k, v in TXT.items()}


def t(key, **kw):
    s = TXT[key][_IDX[LANG]]
    return s.format(**kw) if kw else s


def status_label(code):
    return {UP: t("up"), DOWN: t("down"), FLAT: t("flat")}.get(code, "N/A")


def kw_in(k, text):
    """关键词匹配：词首必须是单词边界；3 个字符以内的短词要求完整单词（避免 'ml' 误中 'html'）。"""
    pat = r"(?<![a-z0-9])" + re.escape(k) + (r"(?![a-z0-9])" if len(k) <= 3 else "")
    return re.search(pat, text) is not None


def clean(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    s = str(v).strip()
    return "" if s.lower() == "nan" else s


# =====================================================================
# 外观：配色 / 字体 / CSS
# =====================================================================
NAVY, BLUE, ORANGE, LINE = "#36425F", "#2F73B8", "#F5A33A", "#D5DCE4"
UP_C, DOWN_C = BLUE, ORANGE   # 上升 = 蓝，下降 = 橙
FONT = '"Barlow","Noto Sans SC","PingFang SC","Microsoft YaHei",system-ui,-apple-system,"Segoe UI",sans-serif'
DISPLAY = '"Oswald","Arial Narrow","Noto Sans SC","PingFang SC","Microsoft YaHei",sans-serif'

CSS = """
:root{--navy:#36425F;--ink:#2B3445;--slate:#5A6890;--band:#E8EDF1;--paper:#F4F6F9;--blue:#2F73B8;--orange:#F5A33A;--line:#D5DCE4}
.stApp{background:#fff!important;color:var(--ink);font-family:__FONT__}
[data-testid="stHeader"]{background:var(--navy)}
[data-testid="stHeader"] *{color:#E8EDF1!important}
.block-container{max-width:1080px;padding-top:4.5rem;padding-bottom:5rem}
#MainMenu,footer{visibility:hidden}
[data-testid="stHeaderActionElements"]{display:none}

h1,h2,h3,h4{font-family:__DISPLAY__!important;color:var(--navy);font-weight:500!important;letter-spacing:.03em}
h2{font-size:1.7rem!important;text-transform:uppercase;margin-top:2.6rem!important;padding-bottom:0!important;border:none!important}
h3{font-size:1.3rem!important;text-transform:uppercase}
hr{margin:2rem 0;border-color:var(--line)}

.hero{background:var(--band);padding:3rem 2.6rem 2.2rem;margin:0 0 1.8rem 0}
.hero h1{font-family:__DISPLAY__!important;font-size:clamp(2.6rem,7vw,4.6rem)!important;font-weight:500!important;text-transform:uppercase;letter-spacing:.02em;color:var(--navy)!important;line-height:1.02;margin:0!important;padding:0!important}
.hero h1::after{content:"";display:inline-block;width:.2em;height:.2em;border-radius:50%;background:var(--orange);margin-left:.12em}
.hero .sub{font-family:__DISPLAY__;text-transform:uppercase;letter-spacing:.08em;color:var(--slate);font-size:1.15rem;margin:.6rem 0 0 0}
.hero .pts{display:grid;grid-template-columns:repeat(2,minmax(0,15rem));gap:.4rem 2.5rem;margin-top:1.6rem;font-family:__DISPLAY__;text-transform:uppercase;letter-spacing:.06em;color:var(--navy);font-size:.95rem}
.hero small{display:block;margin-top:1.8rem;color:var(--slate);font-size:.78rem}

[data-baseweb="tab-list"]{gap:2rem}
button[data-baseweb="tab"]{padding:.5rem 0;background:transparent}
button[data-baseweb="tab"] p{font-family:__DISPLAY__!important;text-transform:uppercase;letter-spacing:.08em;font-size:1rem;color:var(--slate)}
button[data-baseweb="tab"][aria-selected="true"] p{color:var(--navy)}
[data-baseweb="tab-highlight"]{background-color:var(--orange)!important;height:2px!important}
[data-baseweb="tab-border"]{background-color:var(--line)!important;height:1px!important}

[data-testid="stMetric"]{background:transparent;border:none;border-top:1px solid var(--line);border-radius:0;padding:.7rem 0 .2rem}
[data-testid="stMetricLabel"]{color:var(--slate);text-transform:uppercase;letter-spacing:.06em;font-size:.78rem}
[data-testid="stMetricValue"]{font-family:__DISPLAY__;font-size:2rem;font-weight:500;color:var(--navy)}

[class*="st-key-rolecard"]{background:transparent;border:none;border-top:2px solid var(--navy);border-radius:0;padding:1.2rem 0 .4rem;margin:.6rem 0 1.4rem 0}
[class*="st-key-panel"]{background:var(--paper);border:none;border-radius:0;padding:1.6rem 1.8rem .8rem;margin:0 0 1.4rem 0}

[data-testid="stExpander"] details{background:transparent;border:none!important;border-top:1px solid var(--line)!important;border-radius:0}
[data-testid="stAlert"]{border-radius:0;background:var(--paper)!important;border-left:3px solid var(--blue)}
[data-testid="stAlert"] > div{background:transparent!important}
[data-testid="stDataFrame"]{border:1px solid var(--line);border-radius:0;overflow:hidden}

button[kind="primary"],button[data-testid="stBaseButton-primary"]{background:var(--navy);border:1px solid var(--navy);color:#fff;border-radius:2px;padding:.7rem 1.8rem;font-family:__DISPLAY__;text-transform:uppercase;letter-spacing:.1em;font-weight:500}
button[kind="primary"]:hover,button[data-testid="stBaseButton-primary"]:hover{background:var(--blue);border-color:var(--blue);color:#fff}
.stDownloadButton button{border-radius:2px;border:1px solid var(--navy);color:var(--navy);font-family:__DISPLAY__;text-transform:uppercase;letter-spacing:.08em}

[data-testid="stSidebar"]{background:var(--paper);border-right:none}
[data-testid="stSidebar"] .stMarkdown p,[data-testid="stSidebar"] .stMarkdown li{font-size:.86rem;line-height:1.55}

@media (max-width:640px){
 .block-container{padding-left:1rem;padding-right:1rem}
 .hero{padding:1.6rem 1.2rem 1.4rem}
 .hero .pts{grid-template-columns:1fr}
 [class*="st-key-panel"]{padding:1rem}
}
"""


def inject_css(lite):
    web_font = "" if lite else ("@import url('https://fonts.googleapis.com/css2?family=Barlow:wght@400;500;600"
                                "&family=Oswald:wght@400;500&display=swap');")
    css = CSS.replace("__FONT__", FONT).replace("__DISPLAY__", DISPLAY)
    st.markdown(f"<style>{web_font}{css}</style>", unsafe_allow_html=True)


def tidy(chart):
    return chart.configure_view(strokeWidth=0).configure_axis(
        gridColor="#E6EBF0", domainColor="#C7D0DA", tickColor="#C7D0DA",
        labelColor="#40546A", titleColor="#40546A", labelFontSize=12, titleFontSize=12, titleFontWeight="normal")


def box(prefix, idx=None):
    """带 key 的白色卡片容器；旧版 Streamlit 不支持 key 时退回普通带边框容器。"""
    key = f"{prefix}_{idx}" if idx is not None else prefix
    try:
        return st.container(key=key)
    except TypeError:
        return st.container(border=True)


def step_check(label_md, key):
    """勾选框和文字分开放：点文字里的链接只会打开链接，不会误触勾选。"""
    try:
        c1, c2 = st.columns([1, 14], vertical_alignment="center")
    except TypeError:  # 旧版 Streamlit 没有 vertical_alignment
        c1, c2 = st.columns([1, 14])
    with c1:
        ok = st.checkbox("done", key=key, label_visibility="collapsed")
    with c2:
        st.markdown(label_md)
    return ok


# =====================================================================
# AI 趋势信号层（实时来源 + 缓存兜底）
# =====================================================================
def read_cache():
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, encoding="utf-8") as f:
            c = json.load(f)
        return c["items"], c["fetched_at"], True
    return [], None, True


def fetch_live_signals():
    try:
        import feedparser
        ARXIV_API = "https://export.arxiv.org/api/query?search_query=cat:cs.AI&sortBy=submittedDate&sortOrder=descending&max_results=30"
        feed = feedparser.parse(ARXIV_RSS)
        if not feed.entries:
            feed = feedparser.parse(ARXIV_API)
        items = [{"title": e.title, "url": e.link, "published": e.get("published", "N/A")}
                 for e in feed.entries[:30]]
        if not items:
            raise ValueError("empty feed")
        fetched_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump({"fetched_at": fetched_at, "items": items}, f, ensure_ascii=False, indent=2)
        return items, fetched_at, False
    except Exception:
        return read_cache()


@st.cache_data(ttl=3600)
def load_signals(lite):
    # 低流量模式：不联网，只读本地缓存
    return read_cache() if lite else fetch_live_signals()


STOP = {"engineer", "specialist", "manager", "developer", "analyst", "scientist", "senior", "junior", "lead", "and", "the", "for"}
SYNONYMS = {
    "agent": ["agent", "agentic", "autonomous"],
    "prompt": ["prompt", "llm", "language model"],
    "ethic": ["ethic", "safety", "fairness", "bias", "alignment", "trustworthy", "responsible"],
    "risk": ["risk", "safety", "robust"],
    "machine": ["machine learning", "learning"],
    "vision": ["vision", "image", "visual"],
    "data": ["data", "dataset"],
    "nlp": ["language", "nlp", "text"],
    "robot": ["robot", "embodied"],
}


def job_keywords(title):
    words = [w for w in re.split(r"[\s&/,-]+", title.lower()) if len(w) > 2 and w not in STOP and w != "ai"]
    kws = set()
    for w in words:
        kws.add(w)
        for key, syns in SYNONYMS.items():
            if key in w:
                kws.update(syns)
    return kws


def trend_hits(title, signals):
    if not signals:
        return 0
    kws = job_keywords(title)
    if not kws:
        return 0
    return sum(1 for sg in signals if any(k in sg["title"].lower() for k in kws))


# ---- 技能动量：招聘份额变化 + arXiv 研究热度 + 实时招聘占比 ----
HEAT_W = 0.1  # 研究热度的权重（假设值，可调）
SKILL_SYN = {
    "llm": ["llm", "language model"], "prompt": ["prompt"],
    "nlp": ["nlp", "natural language"], "vision": ["vision", "image"],
    "deep learning": ["deep learning", "neural"], "machine learning": ["machine learning"],
    "rag": ["retrieval-augmented", "retrieval augmented", "rag"],
    "agent": ["agent", "agentic"], "reinforcement": ["reinforcement"],
    "ethic": ["ethic", "fairness", "bias", "safety"], "robot": ["robot", "embodied"],
    "generative": ["generative", "diffusion"], "transformer": ["transformer"],
    "multimodal": ["multimodal", "multi-modal"],
}


def _hit(kw, text):
    pat = r"(?<![a-z0-9])" + re.escape(kw)
    if len(kw) <= 4:
        pat += r"(?:s|es)?(?![a-z0-9])"
    return re.search(pat, text) is not None


def skill_heat(skill, signals):
    """某技能在最新 arXiv 论文标题里出现的篇数"""
    if not signals:
        return 0
    s = skill.lower()
    kws = {s}
    for key, syns in SKILL_SYN.items():
        if _hit(key, s):
            kws.update(syns)
    return sum(1 for sg in signals if any(_hit(k, sg["title"].lower()) for k in kws))


def skill_momentum(skill, skill_df, signals):
    """返回 (动量分, 招聘份额变化pp, 论文篇数)"""
    shift = 0.0
    if skill_df is not None:
        hit = skill_df[skill_df["skill"] == skill]
        if len(hit):
            shift = float(hit.iloc[0]["shift_pp"])
    n = len(signals) if signals else 0
    h = skill_heat(skill, signals)
    pct = h / n * 100 if n else 0.0
    # [CHANGED] 加入实时招聘占比
    return shift + HEAT_W * pct + LIVE_MOM_W * LIVE_SHARE.get(skill, 0.0), shift, h


# =====================================================================
# 招聘数据（Kaggle）
# =====================================================================
def find_col(df, candidates):
    for c in candidates:
        if c in df.columns:
            return c
    return None


def split_skills(s):
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return set()
    return {x.strip().lower() for x in re.split(r"[,|;，；]", str(s)) if x.strip()}


def data_period(raw):
    """数据集本身覆盖的招聘时间（不是文件上传日期）"""
    col = find_col(raw, ["posting_date", "date_posted", "posted_date", "date"])
    if col:
        d = pd.to_datetime(raw[col], errors="coerce").dropna()
        if len(d):
            return f"{d.min():%Y-%m-%d} → {d.max():%Y-%m-%d}"
    if "posting_year" in raw.columns:
        ys = raw["posting_year"].dropna().astype(int)
        if len(ys):
            return f"{ys.min()} – {ys.max()}"
    return None


@st.cache_data
def load_data():
    if not os.path.exists(DATA_FILE):
        return None, None, set()
    raw = pd.read_csv(DATA_FILE)
    raw.columns = [c.strip().lower().replace(" ", "_") for c in raw.columns]
    title_col = find_col(raw, ["job_title", "title", "role", "job_role"])
    skill_col = find_col(raw, ["required_skills", "skills", "skill_set", "key_skills"])
    if not title_col or not skill_col:
        return raw, None, set()

    growth_col = find_col(raw, ["demand_growth_yoy_pct", "demand_growth", "growth_yoy"])
    ind_col = find_col(raw, ["industry"])
    rows, vocab = [], Counter()
    for title, g in raw.groupby(title_col):
        counter = Counter()
        for s in g[skill_col].dropna():
            for sk in split_skills(s):
                counter[sk] += 1
                vocab[sk] += 1
        rows.append({
            "job_title": str(title),
            "skills": [s for s, _ in counter.most_common(10)],
            "postings": len(g),
            "growth": float(g[growth_col].mean()) if growth_col else None,
            "industries": list(g[ind_col].value_counts().head(3).index) if ind_col else [],
        })
    jobs = pd.DataFrame(rows)
    jobs["demand_score"] = (jobs["postings"] / jobs["postings"].max() * 100).round(0)
    return raw, jobs, set(vocab.keys())


# =====================================================================
# 未来技能雷达：用数据集 2025 vs 2026 的招聘占比变化判断上升 / 下降
# =====================================================================
def trend_status(v, thr):
    if v >= thr:
        return UP
    if v <= -thr:
        return DOWN
    return FLAT


@st.cache_data
def compute_trends(raw):
    title_col = find_col(raw, ["job_title", "title", "role", "job_role"])
    skill_col = find_col(raw, ["required_skills", "skills", "skill_set", "key_skills"])
    growth_col = find_col(raw, ["demand_growth_yoy_pct", "demand_growth", "growth_yoy"])
    if "posting_year" not in raw.columns or not title_col or not skill_col:
        return None, None, {}
    yr = raw["posting_year"]
    n25, n26 = int((yr == 2025).sum()), int((yr == 2026).sum())
    if n25 == 0 or n26 == 0:
        return None, None, {}
    meta = {"n25": n25, "n26": n26}

    rt = raw.groupby([title_col, "posting_year"]).size().unstack(fill_value=0).reindex(columns=[2025, 2026], fill_value=0)
    role_df = pd.DataFrame({"role": rt.index, "n25": rt[2025].values, "n26": rt[2026].values})
    role_df["share25"] = role_df["n25"] / n25 * 100
    role_df["share26"] = role_df["n26"] / n26 * 100
    role_df["shift_pp"] = role_df["share26"] - role_df["share25"]
    role_df["growth"] = role_df["role"].map(raw.groupby(title_col)[growth_col].mean()) if growth_col else np.nan
    role_df["status"] = role_df["shift_pp"].apply(lambda v: trend_status(v, 0.3))

    g_series = raw[growth_col] if growth_col else pd.Series([np.nan] * len(raw))
    recs = []
    for y, s, g in zip(yr, raw[skill_col], g_series):
        for k in split_skills(s):
            recs.append((k, int(y), g))
    sk = pd.DataFrame(recs, columns=["skill", "year", "growth"])
    cnt = sk.groupby(["skill", "year"]).size().unstack(fill_value=0).reindex(columns=[2025, 2026], fill_value=0)
    skill_df = pd.DataFrame({"skill": cnt.index, "n25": cnt[2025].values, "n26": cnt[2026].values})
    skill_df["share25"] = skill_df["n25"] / n25 * 100
    skill_df["share26"] = skill_df["n26"] / n26 * 100
    skill_df["shift_pp"] = skill_df["share26"] - skill_df["share25"]
    skill_df["growth"] = skill_df["skill"].map(sk.groupby("skill")["growth"].mean())
    skill_df = skill_df[(skill_df["n25"] + skill_df["n26"]) >= 30].reset_index(drop=True)
    skill_df["status"] = skill_df["shift_pp"].apply(lambda v: trend_status(v, 1.0))
    return role_df, skill_df, meta


def diverging_bar(df, cat, val, height=None):
    h = height or max(220, 22 * len(df))
    return alt.Chart(df).mark_bar(cornerRadiusEnd=3).encode(
        x=alt.X(f"{val}:Q", title=t("axis_shift")),
        y=alt.Y(f"{cat}:N", sort="-x", title=None),
        color=alt.condition(alt.datum[val] > 0, alt.value(UP_C), alt.value(DOWN_C)),
        tooltip=[alt.Tooltip(f"{cat}:N"), alt.Tooltip(f"{val}:Q", format="+.2f")],
    ).properties(height=h)


def show_chart(make_chart, table_df):
    """低流量模式下用表格代替图表。"""
    if LITE:
        st.dataframe(table_df, use_container_width=True, hide_index=True)
    else:
        st.altair_chart(tidy(make_chart()), use_container_width=True, theme=None)


def skill_trend_lookup(skill, skill_df):
    if skill_df is None:
        return None
    hit = skill_df[skill_df["skill"] == skill]
    return hit.iloc[0] if len(hit) else None


def render_radar(role_df, skill_df, meta, user_skills, signals=None):
    st.header(t("h_radar"))
    if role_df is None:
        st.info(t("radar_nodata"))
        return
    st.caption(t("radar_method", n25=meta["n25"], n26=meta["n26"]))
    c1, c2 = st.columns(2)
    with c1:
        st.subheader(t("radar_roles"))
        show_chart(lambda: diverging_bar(role_df, "role", "shift_pp"),
                   role_df.sort_values("shift_pp", ascending=False)[["role", "shift_pp"]].round(2))
    with c2:
        st.subheader(t("radar_skills"))
        srt = skill_df.sort_values("shift_pp")
        sk_view = pd.concat([srt.head(10), srt.tail(10)]).drop_duplicates("skill")
        show_chart(lambda: diverging_bar(sk_view, "skill", "shift_pp"),
                   sk_view.sort_values("shift_pp", ascending=False)[["skill", "shift_pp"]].round(2))

    with st.expander(t("radar_exp")):
        t1 = role_df.sort_values("shift_pp", ascending=False)[["role", "status", "shift_pp", "n25", "n26", "growth"]].copy()
        t1["status"] = t1["status"].map(status_label)
        t1.columns = [t("col_role"), t("col_trend"), t("col_shift"), t("col_n25r"), t("col_n26r"), t("col_g_role")]
        st.dataframe(t1.round(2), use_container_width=True, hide_index=True)
        t2 = skill_df.sort_values("shift_pp", ascending=False)[["skill", "status", "shift_pp", "n25", "n26", "growth"]].copy()
        t2["status"] = t2["status"].map(status_label)
        t2.columns = [t("col_skill"), t("col_trend"), t("col_shift"), t("col_n25s"), t("col_n26s"), t("col_g_skill")]
        st.dataframe(t2.round(2), use_container_width=True, hide_index=True)

    if signals:
        rows = []
        for sk in skill_df["skill"]:
            m, sh, h = skill_momentum(sk, skill_df, signals)
            rows.append({t("col_skill"): sk, t("col_shift"): round(sh, 2),
                         t("col_heat"): f"{h}/{len(signals)}", t("col_mom"): round(m, 2)})
        mom = pd.DataFrame(rows).sort_values(t("col_mom"), ascending=False).head(10)
        st.subheader(t("radar_mom_title"))
        st.caption(t("radar_mom_cap", w=HEAT_W))
        st.dataframe(mom, use_container_width=True, hide_index=True)

    if user_skills:
        ups, downs, flats, unknown = [], [], [], []
        for u in sorted(user_skills):
            r = skill_trend_lookup(u, skill_df)
            if r is None:
                unknown.append(u)
            elif r["status"] == UP:
                ups.append(f"{u}（{r['shift_pp']:+.1f}pp）")
            elif r["status"] == DOWN:
                downs.append(f"{u}（{r['shift_pp']:+.1f}pp）")
            else:
                flats.append(u)
        st.markdown(t("your_pos"))
        st.write(t("pos_up", s=', '.join(ups) if ups else t("none")))
        st.write(t("pos_down", s=', '.join(downs) if downs else t("none")))
        st.write(t("pos_flat", s=', '.join(flats) if flats else t("none")))
        if unknown:
            st.caption(t("pos_unknown", s=', '.join(unknown)))


# =====================================================================
# CV 上传 -> 提取文字 -> 识别技能
# =====================================================================
def extract_cv_text(uploaded):
    name = uploaded.name.lower()
    data = uploaded.getvalue()  # 每次从头读，避免 Streamlit 重跑时文件指针已在末尾
    try:
        if name.endswith(".txt") or name.endswith(".md"):
            return data.decode("utf-8", errors="ignore")
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
        if name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for tb in d.tables:  # 很多 CV 把技能放在表格里
                for r in tb.rows:
                    for c in r.cells:
                        parts.append(c.text)
            return "\n".join(parts)
    except Exception as e:
        st.warning(t("cv_fail", e=e))
    return ""


def _norm(s):
    """统一写法：全角转半角、小写、标点换成空格、合并空白（让换行 / 连字符 / 项目符号不影响匹配）。"""
    s = unicodedata.normalize("NFKC", str(s)).lower()
    s = re.sub(r"[^\w+#.]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def detect_skills(text, vocab):
    text = " " + _norm(text) + " "
    found = []
    for sk in vocab:
        n = _norm(sk)
        if len(n) < 3:
            continue
        if re.search(r"(?<![a-z0-9])" + re.escape(n) + r"(?![a-z0-9])", text):
            found.append(sk)
    return sorted(found)


# =====================================================================
# SDG 映射：完整 17 项（中 / English / Bahasa Malaysia）
# =====================================================================
SDG_NAMES = {
    1: ("SDG 1 无贫穷", "SDG 1 No Poverty", "SDG 1 Tiada Kemiskinan"),
    2: ("SDG 2 零饥饿", "SDG 2 Zero Hunger", "SDG 2 Sifar Kelaparan"),
    3: ("SDG 3 良好健康与福祉", "SDG 3 Good Health and Well-being", "SDG 3 Kesihatan dan Kesejahteraan Yang Baik"),
    4: ("SDG 4 优质教育", "SDG 4 Quality Education", "SDG 4 Pendidikan Berkualiti"),
    5: ("SDG 5 性别平等", "SDG 5 Gender Equality", "SDG 5 Kesaksamaan Gender"),
    6: ("SDG 6 清洁饮水和卫生设施", "SDG 6 Clean Water and Sanitation", "SDG 6 Air Bersih dan Sanitasi"),
    7: ("SDG 7 经济适用的清洁能源", "SDG 7 Affordable and Clean Energy", "SDG 7 Tenaga Bersih dan Mampu Milik"),
    8: ("SDG 8 体面工作和经济增长", "SDG 8 Decent Work and Economic Growth", "SDG 8 Pekerjaan Yang Wajar dan Pertumbuhan Ekonomi"),
    9: ("SDG 9 产业、创新和基础设施", "SDG 9 Industry, Innovation and Infrastructure", "SDG 9 Industri, Inovasi dan Infrastruktur"),
    10: ("SDG 10 减少不平等", "SDG 10 Reduced Inequalities", "SDG 10 Mengurangkan Ketidaksamaan"),
    11: ("SDG 11 可持续城市和社区", "SDG 11 Sustainable Cities and Communities", "SDG 11 Bandar dan Komuniti Mampan"),
    12: ("SDG 12 负责任消费和生产", "SDG 12 Responsible Consumption and Production", "SDG 12 Penggunaan dan Pengeluaran Bertanggungjawab"),
    13: ("SDG 13 气候行动", "SDG 13 Climate Action", "SDG 13 Tindakan Iklim"),
    14: ("SDG 14 水下生物", "SDG 14 Life Below Water", "SDG 14 Kehidupan Bawah Air"),
    15: ("SDG 15 陆地生物", "SDG 15 Life on Land", "SDG 15 Kehidupan di Darat"),
    16: ("SDG 16 和平、正义与强大机构", "SDG 16 Peace, Justice and Strong Institutions", "SDG 16 Keamanan, Keadilan dan Institusi yang Kukuh"),
    17: ("SDG 17 促进目标实现的伙伴关系", "SDG 17 Partnerships for the Goals", "SDG 17 Perkongsian untuk Mencapai Matlamat"),
}


def sdg_name(n):
    return SDG_NAMES[n][_IDX[LANG]]


# 岗位规则：按顺序匹配第一条 (关键词, SDG, (中, en, ms 说明))
ROLE_RULES = [
    (["ethic"], [16, 10], ("让 AI 公平、透明、负责任，减少算法带来的不平等",
                           "Making AI fair, transparent and accountable, reducing algorithmic inequality",
                           "Menjadikan AI adil, telus dan bertanggungjawab serta mengurangkan ketidaksamaan algoritma")),
    (["compliance", "governance", "policy"], [16, 8], ("建立 AI 合规与治理制度，保障机构可信与企业稳健经营",
                                                      "Building AI compliance and governance so institutions are trustworthy and businesses run soundly",
                                                      "Membina pematuhan dan tadbir urus AI untuk institusi yang dipercayai dan perniagaan yang stabil")),
    (["security"], [16, 9], ("保护 AI 系统与数据安全，维护数字基础设施的可信度",
                             "Protecting AI systems and data to keep digital infrastructure trustworthy",
                             "Melindungi sistem AI dan data untuk memastikan infrastruktur digital dipercayai")),
    (["research"], [9, 4], ("推动 AI 前沿创新，并促进知识传播",
                            "Advancing frontier AI innovation and spreading knowledge",
                            "Memajukan inovasi AI termaju dan menyebarkan pengetahuan")),
    (["product", "business"], [8, 9], ("把 AI 转化为企业效率与高质量就业",
                                       "Turning AI into business efficiency and quality jobs",
                                       "Menukar AI kepada kecekapan perniagaan dan pekerjaan berkualiti")),
    (["infrastructure", "mlops"], [9, 13], ("构建稳定、可扩展、更节能的 AI 算力与部署基础设施",
                                           "Building stable, scalable and more energy-efficient AI compute and deployment infrastructure",
                                           "Membina infrastruktur pengkomputeran dan penempatan AI yang stabil, berskala dan lebih jimat tenaga")),
    (["robot"], [9, 8], ("推动智能自动化与先进制造", "Driving intelligent automation and advanced manufacturing",
                         "Memacu automasi pintar dan pembuatan termaju")),
    (["vision", "multimodal"], [9, 3], ("构建视觉与多模态 AI，可应用于医疗影像等场景",
                                        "Building vision and multimodal AI, e.g. for medical imaging",
                                        "Membina AI penglihatan dan multimodal, contohnya untuk pengimejan perubatan")),
    (["data scientist", "data engineer"], [9, 8], ("用数据驱动决策，支撑产业与企业增长",
                                                   "Using data to drive decisions and support industry and business growth",
                                                   "Menggunakan data untuk memacu keputusan dan menyokong pertumbuhan industri dan perniagaan")),
    (["nlp", "llm", "rag", "prompt", "generative", "agent"], [9, 4], ("构建语言类 AI 应用，可用于教育与知识获取",
                                                                      "Building language-based AI applications for education and access to knowledge",
                                                                      "Membina aplikasi AI berasaskan bahasa untuk pendidikan dan akses kepada pengetahuan")),
    (["engineer", "developer", "architect", "scientist"], [9], ("构建 AI 技术与数字基础设施",
                                                               "Building AI technology and digital infrastructure",
                                                               "Membina teknologi AI dan infrastruktur digital")),
]
DEFAULT_ROLE_SDG = ([8, 9], ("推动高质量就业与产业创新", "Driving quality jobs and industry innovation",
                             "Memacu pekerjaan berkualiti dan inovasi industri"))

SKILL_RULES = [
    (["ethic", "fairness", "bias"], [10, 16]),
    (["policy", "legal", "eu ai act", "compliance", "audit", "governance", "risk", "red teaming", "security", "cyber", "privacy"], [16]),
    (["ai literacy", "publishing", "teach", "training", "mentor"], [4]),
    (["communication", "leadership", "stakeholder", "agile", "problem solving", "business", "strategy", "management",
      "excel", "erp", "accounting", "financ"], [8]),
    (["research"], [9, 4]),
    (["simulation", "control systems", "ros", "opencv", "image processing", "vision", "planning systems", "robot"], [9]),
]
DEFAULT_SKILL_SDG = [9]

INDUSTRY_SDG = {
    "Healthcare": [3], "Education": [4], "Energy": [7, 13], "Government": [16], "Finance": [8],
    "Manufacturing": [9, 12], "Automotive": [9, 11], "Retail": [12], "Media": [16],
    "Research": [9, 4], "Consulting": [8, 17], "Technology": [9],
}


def role_sdg(title):
    tl = title.lower()
    for kws, sdgs, desc in ROLE_RULES:
        if any(kw_in(k, tl) for k in kws):
            return sdgs, desc[_IDX[LANG]]
    return DEFAULT_ROLE_SDG[0], DEFAULT_ROLE_SDG[1][_IDX[LANG]]


def skill_sdg(skill):
    s = skill.lower()
    for kws, sdgs in SKILL_RULES:
        if any(kw_in(k, s) for k in kws):
            return sdgs
    return DEFAULT_SKILL_SDG


def sdg_label(nums):
    return " & ".join(sdg_name(n) for n in nums)


def tag_skills(skills):
    return ", ".join(f"{s} 〔SDG {'/'.join(map(str, skill_sdg(s)))}〕" for s in skills)


def industry_sdg_text(industries):
    parts = [f"{i}（SDG {'/'.join(map(str, INDUSTRY_SDG[i]))}）" for i in industries if i in INDUSTRY_SDG]
    return "、".join(parts)


def render_sdg_coverage(jobs_df, vocab):
    with st.expander(t("h_cov"), expanded=True):
        role_cnt, skill_cnt, ind_cnt = Counter(), Counter(), Counter()
        for title in jobs_df["job_title"]:
            role_cnt.update(role_sdg(title)[0])
        for s in vocab:
            skill_cnt.update(skill_sdg(s))
        for sdgs in INDUSTRY_SDG.values():
            ind_cnt.update(sdgs)
        rows = []
        for n in SDG_NAMES:
            direct = role_cnt[n] + skill_cnt[n]
            rows.append({"SDG": sdg_name(n), t("c_roles"): role_cnt[n], t("c_skills"): skill_cnt[n],
                         t("c_ind"): ind_cnt[n],
                         t("c_cov"): t("cov_direct") if direct else (t("cov_ind") if ind_cnt[n] else t("cov_none"))})
        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
        st.caption(t("cov_cap"))


def render_region_sdg(pairs):
    """pairs: [(region, role_title), ...] —— 按地区汇总 SDG 贡献"""
    st.subheader(t("h_region"))
    if not any(r for r, _ in pairs):
        st.caption(t("region_hint"))
        return
    people, sdgs = Counter(), {}
    for region, title in pairs:
        region = region or t("no_region")
        people[region] += 1
        sdgs.setdefault(region, Counter()).update(role_sdg(title)[0])
    rows = []
    for region, n in people.most_common():
        top = "；".join(f"{sdg_name(k)}（{v}）" for k, v in sdgs[region].most_common(3))
        rows.append({t("region"): region, t("n_people"): n, t("top_sdg"): top})
    st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
    st.caption(t("region_cap"))


# =====================================================================
# 免费/低成本学习资源
# =====================================================================
RESOURCES = [
    (["python"], "Kaggle Learn · Python", "https://www.kaggle.com/learn/python", "c_free", 5),
    (["pandas", "data analysis", "analytics", "data"], "Kaggle Learn · Pandas", "https://www.kaggle.com/learn/pandas", "c_free", 4),
    (["sql", "database"], "Kaggle Learn · Intro to SQL", "https://www.kaggle.com/learn/intro-to-sql", "c_free", 4),
    (["machine learning", "ml", "deep learning", "tensorflow", "pytorch", "model"], "Google · Machine Learning Crash Course",
     "https://developers.google.com/machine-learning/crash-course", "c_free", 15),
    (["llm", "prompt", "generative", "nlp", "language"], "DeepLearning.AI · Short Courses", "https://www.deeplearning.ai/short-courses/", "c_mostly", 3),
    (["ethic", "risk", "govern", "compliance", "privacy", "policy", "safety"], "Elements of AI (includes AI ethics)", "https://www.elementsofai.com/", "c_free", 30),
    (["product", "strategy", "management", "communication"], "Coursera · AI For Everyone", "https://www.coursera.org/learn/ai-for-everyone", "c_audit", 6),
]


def pick_resource(skill):
    s = skill.lower()
    for kws, name, url, cost, hrs in RESOURCES:
        if any(kw_in(k, s) for k in kws):
            return name, url, cost, hrs
    return None


def no_resource_skills(skills):
    """没有匹配资源的技能：如实列出，不编造课程。"""
    return [s for s in skills if pick_resource(s) is None]


def build_plan(missing, weekly_hours, skill_df=None, signals=None):
    """missing 按动量排序：招聘份额上升 + arXiv 研究热度高的技能优先。
    找不到资源的技能不会被硬配课程；steps 可能为空。"""
    order = list(missing)
    order.sort(key=lambda s: -skill_momentum(s, skill_df, signals)[0])
    steps, seen = [], set()
    for sk in order:
        r = pick_resource(sk)
        if r and r[0] not in seen:
            seen.add(r[0])
            steps.append((sk, *r))
        if len(steps) == 3:
            break
    learn_hours = sum(s[4] for s in steps)
    project_hours, cv_hours = 8, 3
    total = learn_hours + project_hours + cv_hours
    weeks = max(1, math.ceil(total / max(1, weekly_hours)))
    return steps, total, weeks, project_hours, cv_hours


# =====================================================================
# 匹配引擎（个人 / 团队共用）
# =====================================================================
SDG_W = 5  # 每个与用户所选 SDG 重合的项，在排序中加几分（假设值，只影响排序）

BG_MAP = {"account": ["financ", "analyst", "risk", "audit"], "financ": ["financ", "analyst", "risk"],
          "engineer": ["engineer", "developer", "ml", "machine"], "computer": ["engineer", "developer", "data", "machine"],
          "market": ["product", "analyst", "strategy"], "business": ["product", "strategy", "analyst"],
          "law": ["ethic", "risk", "govern", "compliance"], "student": []}
FAIR_BACKGROUNDS = ["accounting", "marketing", "engineer", "law", "computer", "other"]

# 背景下拉选项：键（英文）要包含 BG_MAP 的关键词；显示文字随语言切换 (中, English, Bahasa Malaysia)
ROLE_CHOICES = {
    "student of accounting and finance": ("会计与金融专业学生", "Student of accounting and finance", "Pelajar perakaunan dan kewangan"),
    "accountant / finance professional": ("会计 / 金融从业者", "Accountant / finance professional", "Akauntan / profesional kewangan"),
    "student of computer science": ("计算机专业学生", "Student of computer science", "Pelajar sains komputer"),
    "engineer / developer": ("工程师 / 开发者", "Engineer / developer", "Jurutera / pembangun"),
    "marketing professional": ("市场营销从业者", "Marketing professional", "Profesional pemasaran"),
    "business / management student": ("商科 / 管理专业学生", "Business / management student", "Pelajar perniagaan / pengurusan"),
    "law / compliance professional": ("法律 / 合规从业者", "Law / compliance professional", "Profesional undang-undang / pematuhan"),
    "other": ("其他", "Other", "Lain-lain"),
}


def skill_match(u, s):
    if u == s:
        return True
    for a, b in ((u, s), (s, u)):
        if len(a) >= 3 and re.search(r"(?<![a-z0-9])" + re.escape(a) + r"(?![a-z0-9])", b):
            return True
    return False


def match_roles(user_skills, user_role, jobs_df, signals=None, role_shift=None, bg_w=10, pref_sdgs=None):
    """匹配度 = (100-bg_w)% × 技能覆盖率 + bg_w% × 背景相关度。
    趋势加权（雷达份额变化 + arXiv 热度 + 实时招聘需求）和 SDG 偏好加权只进入 rank_score 用于排序，不改变匹配度分数。"""
    results = []
    role_text = (user_role or "").lower()
    role_shift = role_shift or {}
    for _, row in jobs_df.iterrows():
        req = row["skills"]
        matched = [s for s in req if any(skill_match(u, s) for u in user_skills)]
        missing = [s for s in req if s not in matched]
        coverage = len(matched) / len(req) if req else 0
        bg = 0.0
        for key, targets in BG_MAP.items():
            if key in role_text and any(tt in row["job_title"].lower() for tt in targets):
                bg = 1.0
        score = round((100 - bg_w) * coverage + bg_w * bg)
        hits = trend_hits(row["job_title"], signals)
        trend_b = (float(np.clip(role_shift.get(row["job_title"], 0.0) * 3, -3, 3)) + min(hits, 3) * 0.5
                   + min(sum(1 for s in req if skill_heat(s, signals) > 0), 4) * 0.5)
        # [NEW] 实时招聘需求：所需技能中有多少在实时广告里占比 ≥ LIVE_MIN（最多计 4 个）
        trend_b += LIVE_W * min(sum(1 for s in req if LIVE_SHARE.get(s, 0.0) >= LIVE_MIN), 4)
        sdg_hit = sorted(set(role_sdg(row["job_title"])[0]) & set(pref_sdgs or []))
        sdg_b = SDG_W * len(sdg_hit)
        results.append({**row.to_dict(), "matched": matched, "missing": missing, "coverage": coverage,
                        "bg": bg, "match_score": score, "trend_hits": hits,
                        "trend_bonus": trend_b, "sdg_bonus": sdg_b, "sdg_hit": sdg_hit,
                        "rank_score": score + trend_b + sdg_b})
    return pd.DataFrame(results).sort_values(["rank_score", "match_score", "postings"], ascending=False)


def base_top(res, n=3):
    """不含趋势信号的排序（但保留用户的 SDG 偏好），用来对比“重新规划”前后的变化。"""
    tmp = res.assign(_base=res["match_score"] + res["sdg_bonus"])
    return tmp.sort_values(["_base", "match_score", "postings"], ascending=False)["job_title"].head(n).tolist()


def render_fairness(user_skills, jobs_df, bg_w):
    with st.expander(t("h_fair"), expanded=True):
        if not user_skills:
            st.info(t("fair_need"))
            return
        st.caption(t("fair_cap", bw=bg_w))
        ref = match_roles(user_skills, "other", jobs_df, None, None, bg_w)
        ref_role = base_top(ref, 1)[0]
        rows, scores = [], []
        for bgname in FAIR_BACKGROUNDS:
            res = match_roles(user_skills, bgname, jobs_df, None, None, bg_w)
            sc = int(res[res["job_title"] == ref_role].iloc[0]["match_score"])
            scores.append(sc)
            rows.append({t("fc_bg"): bgname, t("fc_score"): sc, t("fc_top"): base_top(res, 1)[0]})
        st.write(t("fair_role", r=ref_role))
        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
        gap_pts = max(scores) - min(scores)
        if bg_w == 0:
            st.success(t("fair_zero"))
        else:
            st.success(t("fair_result", g=gap_pts, bw=bg_w))
        st.caption(t("fair_limit"))


# =====================================================================
# 团队 / 批量模式（HR·L&D、TVET·大学职业中心）
# =====================================================================
TEMPLATE_CSV = ("name,background,skills,region\nMember A,accounting,Excel; ERP; Accounting,Selangor\n"
                "Member B,marketing,Communication; Excel; Product Strategy,Penang\n")


def collect_members():
    st.markdown(t("enter_members"))
    st.download_button(t("dl_tpl"), TEMPLATE_CSV, file_name="team_template.csv", mime="text/csv")
    up = st.file_uploader(t("up_csv"), type=["csv"], key="team_csv")
    if up is not None:
        try:
            df = pd.read_csv(up)
            df.columns = [c.strip().lower() for c in df.columns]
            n_col = find_col(df, ["name", "姓名"])
            b_col = find_col(df, ["background", "role", "major", "专业", "背景"])
            s_col = find_col(df, ["skills", "skill", "技能"])
            r_col = find_col(df, ["region", "state", "地区", "wilayah"])
            if not s_col:
                st.error(t("csv_needs"))
                return []
            members = []
            for i, r in df.iterrows():
                members.append({"name": clean(r[n_col]) if n_col else t("member", i=i + 1),
                                "background": clean(r[b_col]) if b_col else "",
                                "region": clean(r[r_col]) if r_col else "",
                                "skills": split_skills(r[s_col])})
            return members
        except Exception as e:
            st.error(t("csv_fail", e=e))
            return []
    default = pd.DataFrame({
        "name": ["Member A", "Member B", "Member C", "Member D"],
        "background": ["accounting", "accounting", "marketing", "engineer"],
        "skills": ["Excel, ERP, Accounting", "Excel, SQL, Financial Analysis",
                   "Communication, Excel, Product Strategy", "Python, SQL, Docker, Git"],
        "region": ["Selangor", "Selangor", "Penang", "Johor"],
    })
    edited = st.data_editor(default, num_rows="dynamic", use_container_width=True, key="team_editor")
    members = []
    for i, r in edited.iterrows():
        if clean(r.get("skills")):
            members.append({"name": clean(r.get("name")) or t("member", i=i + 1),
                            "background": clean(r.get("background")),
                            "region": clean(r.get("region")),
                            "skills": split_skills(r["skills"])})
    st.caption(t("mem_cap"))
    return members


def heatmap_chart(long_df, order):
    return alt.Chart(long_df).mark_rect(stroke="white").encode(
        x=alt.X("skill:N", sort=order, title=None, axis=alt.Axis(labelAngle=-40)),
        y=alt.Y("member:N", title=None),
        color=alt.Color("state:N", scale=alt.Scale(domain=[t("have"), t("gap")], range=[UP_C, DOWN_C]),
                        legend=alt.Legend(title=None)),
        tooltip=["member", "skill", "state"],
    ).properties(height=max(120, 28 * long_df["member"].nunique()))


def render_team_mode(jobs_df, role_df, skill_df, signals, bg_w):
    st.header(t("h_team"))
    st.info(t("team_human"))
    opt_hr, opt_tvet = t("aud_hr"), t("aud_tvet")
    audience = st.radio(t("aud_label"), [opt_hr, opt_tvet], horizontal=True, key="audience")
    hours = st.slider(t("team_hours"), 2, 20, 5, key="team_hours")
    members = collect_members()
    if not members:
        st.info(t("need_member"))
        return

    trend_map = dict(zip(role_df["role"], role_df["status"])) if role_df is not None else {}
    shift_map = dict(zip(role_df["role"], role_df["shift_pp"])) if role_df is not None else {}

    per_member = []
    for m in members:
        res = match_roles(m["skills"], m["background"], jobs_df, signals, shift_map, bg_w)
        per_member.append((m, res))

    if audience == opt_hr:
        # ---- HR：针对一个目标岗位做团队技能缺口分析 ----
        top1 = Counter(res.iloc[0]["job_title"] for _, res in per_member)
        titles = sorted(jobs_df["job_title"])
        default_title = top1.most_common(1)[0][0]
        target = st.selectbox(t("target_sel"), titles, index=titles.index(default_title))
        row = jobs_df[jobs_df["job_title"] == target].iloc[0]
        req = row["skills"]

        cov, member_rows, heat_rows, total_hours = [], [], [], 0
        miss_counter = Counter()
        for m, res in per_member:
            r = res[res["job_title"] == target].iloc[0]
            for s in r["missing"]:
                miss_counter[s] += 1
            _, tot, wk, _, _ = build_plan(r["missing"], hours, skill_df, signals)
            total_hours += tot
            member_rows.append({t("col_member"): m["name"], t("col_bg"): m["background"], t("match_pct"): r["match_score"],
                                t("n_gap"): len(r["missing"]), t("est_h"): tot, t("est_wk"): wk})
            for s in req:
                heat_rows.append({"member": m["name"], "skill": s,
                                  "state": t("have") if s in r["matched"] else t("gap")})
        n = len(per_member)
        for s in req:
            have = sum(1 for _, res in per_member if s in res[res["job_title"] == target].iloc[0]["matched"])
            cov.append({"skill": s, "coverage": round(have / n * 100)})

        c1, c2, c3 = st.columns(3)
        c1.metric(t("m_team_avg"), f"{np.mean([r[t('match_pct')] for r in member_rows]):.0f}%")
        c2.metric(t("m_team_hours"), t("hours_unit", h=total_hours))
        c3.metric(t("m_target_sig"), status_label(trend_map.get(target)),
                  help=t("sig_help") + (f"（{shift_map[target]:+.2f}pp）" if target in shift_map else ""))

        st.subheader(t("team_cov_title", t=target))
        cov_df = pd.DataFrame(cov)
        show_chart(lambda: alt.Chart(cov_df).mark_bar(color=NAVY, cornerRadiusEnd=3).encode(
            x=alt.X("coverage:Q", title=t("team_cov_axis"), scale=alt.Scale(domain=[0, 100])),
            y=alt.Y("skill:N", sort="-x", title=None),
            tooltip=["skill", "coverage"]).properties(height=max(200, 26 * len(cov_df))), cov_df)

        # ---- 技能缺口热力图（成员 × 技能） ----
        st.subheader(t("heat_title"))
        heat_df = pd.DataFrame(heat_rows)
        order = [c["skill"] for c in sorted(cov, key=lambda c: -c["coverage"])]
        if LITE:
            pv = heat_df.pivot(index="member", columns="skill", values="state").reindex(columns=order)
            st.dataframe(pv.replace({t("have"): "✅", t("gap"): "❌"}), use_container_width=True)
        else:
            st.altair_chart(tidy(heatmap_chart(heat_df, order)), use_container_width=True, theme=None)

        st.subheader(t("prio_title"))
        gap_rows = []
        for s, c in miss_counter.most_common():
            r = pick_resource(s)
            gap_rows.append({t("col_skill"): s, t("n_missing"): c, t("pct"): round(c / n * 100),
                             t("arrange"): t("group") if c / n >= 0.5 else t("indiv"),
                             t("res"): r[0] if r else t("no_res"), t("cost"): t(r[2]) if r else "",
                             "SDG": "/".join(map(str, skill_sdg(s)))})
        st.dataframe(pd.DataFrame(gap_rows), use_container_width=True, hide_index=True)
        st.subheader(t("member_detail"))
        st.dataframe(pd.DataFrame(member_rows), use_container_width=True, hide_index=True)
        sdgs, desc = role_sdg(target)
        st.info(t("sdg_chain_team", title=target, sdg=sdg_label(sdgs), desc=desc))
        render_region_sdg([(m["region"], target) for m, _ in per_member])
        st.markdown(t("hrd"))
        st.caption(t("team_cap"))

    else:
        # ---- 职业中心：学生群体画像 ----
        rows, top1_counter, gap_counter, sdg_counter, pairs = [], Counter(), Counter(), Counter(), []
        for m, res in per_member:
            top3 = res.head(3)
            t1 = top3.iloc[0]
            top1_counter[t1["job_title"]] += 1
            gap_counter.update(t1["missing"])
            sdg_counter.update(role_sdg(t1["job_title"])[0])
            pairs.append((m["region"], t1["job_title"]))
            rows.append({
                t("student"): m["name"], t("col_bg"): m["background"],
                t("top_role"): t1["job_title"], t("match_col"): t1["match_score"],
                t("col_trend"): status_label(trend_map.get(t1["job_title"])),
                t("alts"): " / ".join(f"{r['job_title']}({r['match_score']}%)" for _, r in top3.iloc[1:].iterrows()),
                t("top_gap"): ", ".join(t1["missing"][:3]),
            })
        out = pd.DataFrame(rows)
        c1, c2 = st.columns(2)
        c1.metric(t("n_students"), len(out))
        c2.metric(t("avg_match"), f"{out[t('match_col')].mean():.0f}%")

        a, b = st.columns(2)
        with a:
            st.subheader(t("top_dist"))
            d1 = pd.DataFrame(top1_counter.most_common(), columns=["role", "count"])
            show_chart(lambda: alt.Chart(d1).mark_bar(color=NAVY, cornerRadiusEnd=3).encode(
                x="count:Q", y=alt.Y("role:N", sort="-x", title=None), tooltip=["role", "count"]), d1)
        with b:
            st.subheader(t("common_gaps"))
            d2 = pd.DataFrame(gap_counter.most_common(10), columns=["skill", "count"])
            show_chart(lambda: alt.Chart(d2).mark_bar(color=NAVY, cornerRadiusEnd=3).encode(
                x="count:Q", y=alt.Y("skill:N", sort="-x", title=None), tooltip=["skill", "count"]), d2)
        st.subheader(t("per_student"))
        st.dataframe(out, use_container_width=True, hide_index=True)
        st.download_button(t("dl_result"), out.to_csv(index=False).encode("utf-8-sig"),
                           file_name="cohort_recommendations.csv", mime="text/csv")
        top_sdg = ", ".join(sdg_name(n) for n, _ in sdg_counter.most_common(3))
        st.success(t("cohort_sdg", s=top_sdg))
        render_region_sdg(pairs)
        st.caption(t("curriculum"))


# =====================================================================
# [NEW] 实时招聘信号（Arbeitnow 公开 API；失败或低流量模式时用本地缓存）
# =====================================================================
def read_live_cache():
    if os.path.exists(LIVE_CACHE):
        try:
            with open(LIVE_CACHE, encoding="utf-8") as f:
                c = json.load(f)
            return c["items"], c["meta"], True
        except Exception:
            pass
    return [], {}, True


def fetch_live_jobs():
    import urllib.request
    try:
        items, stamps = [], []
        for page in (1, 2, 3):
            req = urllib.request.Request(f"{LIVE_JOBS_URL}?page={page}", headers={"User-Agent": "Aiversecity-demo/1.0"})
            with urllib.request.urlopen(req, timeout=10) as r:
                data = json.loads(r.read().decode("utf-8")).get("data", [])
            if not data:
                break
            for j in data:
                body = html.unescape(re.sub(r"<[^>]+>", " ", j.get("description") or ""))
                text = " ".join([j.get("title") or "", " ".join(j.get("tags") or []), body[:1200]]).lower()
                items.append({"title": j.get("title") or "", "url": j.get("url") or "", "text": text})
                if j.get("created_at"):
                    stamps.append(int(j["created_at"]))
        if not items:
            raise ValueError("empty")

        def day(ts):
            return datetime.datetime.fromtimestamp(ts).strftime("%Y-%m-%d")

        meta = {"fetched_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
                "oldest": day(min(stamps)) if stamps else "?", "newest": day(max(stamps)) if stamps else "?"}
        with open(LIVE_CACHE, "w", encoding="utf-8") as f:
            json.dump({"meta": meta, "items": items}, f, ensure_ascii=False)
        return items, meta, False
    except Exception:
        return read_live_cache()


@st.cache_data(ttl=3600)
def load_live_jobs(lite):
    return read_live_cache() if lite else fetch_live_jobs()


@st.cache_data(ttl=3600)
def live_skill_share(items, vocab_t):
    """每个技能被多少 % 的实时广告提到（至少 3 条才计入，避免噪音）。"""
    if not items:
        return {}, 0
    cnt = Counter()
    for it in items:
        cnt.update(set(detect_skills(it["text"], vocab_t)))
    n = len(items)
    return {k: v / n * 100 for k, v in cnt.items() if v >= 3}, n


def render_live_panel(skill_df):
    st.subheader(t("live_title"))
    if not LIVE_N:
        st.info(t("live_none"))
        return
    st.markdown(t("live_src"))
    st.caption(t("live_cap", w=LIVE_MOM_W))
    st.write(t("live_n", n=LIVE_N))
    st.write(t("live_when", ts=LIVE_META.get("fetched_at", "N/A")) + (t("cached") if LIVE_CACHED else ""))
    st.write(t("live_range", r=f"{LIVE_META.get('oldest', '?')} → {LIVE_META.get('newest', '?')}"))
    base = dict(zip(skill_df["skill"], skill_df["share26"])) if skill_df is not None else {}
    rows = []
    for s, v in sorted(LIVE_SHARE.items(), key=lambda kv: -kv[1])[:10]:
        b = base.get(s)
        rows.append({t("col_skill"): s, t("live_col_live"): round(v, 1),
                     t("live_col_ds"): None if b is None else round(b, 1),
                     t("live_col_diff"): None if b is None else round(v - b, 1)})
    st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
    st.caption(t("live_bias"))


# =====================================================================
# 页面开始（侧边栏：语言、低流量模式放最前面）
# =====================================================================
lang_name = st.sidebar.selectbox("Language / 语言 / Bahasa", list(LANGS), index=0, key="lang")
LANG = LANGS[lang_name]
LITE = st.sidebar.checkbox(t("lite"), key="lite")
inject_css(LITE)

raw_df, jobs_df, vocab = load_data()
signals, fetched_at, from_cache = load_signals(LITE)
# [NEW] 实时招聘信号
live_items, LIVE_META, LIVE_CACHED = load_live_jobs(LITE)
LIVE_SHARE, LIVE_N = live_skill_share(live_items, tuple(sorted(vocab))) if vocab else ({}, 0)
role_df, skill_df, trend_meta = compute_trends(raw_df) if raw_df is not None else (None, None, {})
role_shift = dict(zip(role_df["role"], role_df["shift_pp"])) if role_df is not None else {}

# ---------- 侧边栏：AI 披露常驻，其余收进折叠面板 ----------
with st.sidebar.container(border=True):
    st.markdown(f"**{t('sb_disc_title')}**")
    st.markdown(t("sb_disc"))

with st.sidebar.expander(t("sb_signals")):
    st.markdown(f"**{t('sb_src1')}**")
    if raw_df is None:
        st.error(t("sb_missing", f=DATA_FILE))
    else:
        mtime = datetime.date.fromtimestamp(os.path.getmtime(DATA_FILE)).isoformat()
        st.markdown(t("sb_source", url=KAGGLE_URL))
        st.write(t("sb_filedate", d=mtime))
        period = data_period(raw_df)
        if period:
            st.write(t("sb_period", r=period))
        st.write(t("sb_samples", n=len(raw_df)))
    st.markdown(f"**{t('sb_src2')}**")
    if signals:
        for sg in signals[:5]:
            st.markdown(f"- [{sg['title'][:70]}]({sg['url']})  \n  {t('sb_pub', p=sg['published'])}")
        st.caption(t("sb_arxiv", ts=fetched_at) + (t("cached") if from_cache else ""))
    else:
        st.warning(t("sb_nosig"))

with st.sidebar.expander(t("sb_assump")):
    skills_only = st.checkbox(t("skills_only"), key="skills_only")
    mult_slow = st.number_input(t("mult_slow"), 0.1, 3.0, 0.5, 0.1)
    mult_mod = st.number_input(t("mult_mod"), 0.1, 3.0, 1.0, 0.1)
    mult_fast = st.number_input(t("mult_fast"), 0.1, 5.0, 1.5, 0.1)
    spread = st.slider(t("spread"), 10, 60, 30)
BG_W = 0 if skills_only else 10

# ---------- 顶部 hero ----------
_pts = "".join(f"<span>{html.escape(x)}</span>" for x in t("hero_pts").split("|"))
st.markdown(
    f'<div class="hero"><h1>Aiversecity</h1><p class="sub">{html.escape(t("app_sub"))}</p>'
    f'<div class="pts">{_pts}</div>'
    f'<small>Future-Ready Workforce Domain | Empowering Humans in the AI Era</small></div>',
    unsafe_allow_html=True)
st.info(t("human_note"))

if raw_df is None:
    st.error(t("err_nodata", f=DATA_FILE))
    st.stop()
if jobs_df is None:
    st.error(t("err_cols"))
    st.write(list(raw_df.columns))
    st.stop()

tab_me, tab_team = st.tabs([t("tab_me"), t("tab_team")])

with tab_me:
    sub_path, sub_trend, sub_trust = st.tabs([t("sub_path"), t("sub_trend"), t("sub_trust")])

    # ======================= 我的路径 =======================
    with sub_path:
        st.header(t("h_intake"))
        with box("panel", "intake"):
            col1, col2 = st.columns(2)
            with col1:
                # 下拉列表：内部值为英文键（含 BG_MAP 关键词），显示文字随语言切换
                user_role = st.selectbox(
                    t("in_role"),
                    list(ROLE_CHOICES),
                    index=0,
                    format_func=lambda k: ROLE_CHOICES[k][_IDX[LANG]],
                    key="user_role",
                )
                user_skills_input = st.text_area(t("in_skills"), "Excel, ERP, Accounting, Financial Analysis")
                cv_file = st.file_uploader(t("in_cv"), type=["pdf", "docx", "txt", "md"])
                st.caption(t("cv_notice"))
            with col2:
                st.markdown(t("q_title"))
                weekly_hours = st.slider(t("q_hours"), 2, 20, 5)
                learn_pref = st.selectbox(t("q_pref"), [t("pref_self"), t("pref_proj"), t("pref_course")])
                mind_opts = [t("mind_1"), t("mind_2"), t("mind_3")]
                growth_mindset = st.select_slider(t("q_mind"), mind_opts, value=mind_opts[1])
                pref_sdgs = st.multiselect(t("q_sdg"), list(SDG_NAMES), format_func=sdg_name, key="pref_sdgs")

        user_skills = split_skills(user_skills_input)
        if cv_file is not None:
            cv_text = extract_cv_text(cv_file)
            if not cv_text.strip():
                st.warning(t("cv_empty"))
            else:
                st.caption(t("cv_chars", n=len(cv_text)))
                cv_found = detect_skills(cv_text, vocab)
                if cv_found:
                    st.success(t("cv_found", n=len(cv_found), s=', '.join(cv_found[:15])))
                    user_skills |= set(cv_found)
                else:
                    st.info(t("cv_none"))

        st.markdown("---")

        # 点击一次后用 session_state 记住，之后切换岗位 / 勾选进度都不会让结果消失
        if st.button(t("btn_run"), type="primary"):
            st.session_state["run"] = True

        if st.session_state.get("run"):
            st.header(t("h_match"))
            st.caption(t("match_cap", cw=100 - BG_W, bw=BG_W))

            res_all = match_roles(user_skills, user_role, jobs_df, signals, role_shift, BG_W, pref_sdgs)
            top = res_all.head(3)
            if pref_sdgs:
                st.caption(t("sdg_pref_on", w=SDG_W))

            # 自动重新规划：对比"不含趋势信号"和"含趋势信号"的推荐
            before, after = base_top(res_all), top["job_title"].tolist()
            if before != after:
                st.info(t("replan_changed", old=" → ".join(before), new=" → ".join(after)))
            else:
                st.caption(t("replan_same"))

            trend_map_me = dict(zip(role_df["role"], role_df["status"])) if role_df is not None else {}
            chosen_sdgs = Counter()
            for tt in top["job_title"]:
                chosen_sdgs.update(role_sdg(tt)[0])

            # 一次只看一个岗位：前三名用单选切换
            labels = [f"{r['job_title']}  ({r['match_score']}%)" for _, r in top.iterrows()]
            pick = st.radio(t("pick_role"), labels, horizontal=True, key="role_pick")
            rank = labels.index(pick) if pick in labels else 0
            row = top.iloc[rank]
            title = row["job_title"]
            sdgs, sdg_desc = role_sdg(title)

            with box("rolecard", rank):
                st.subheader(t("rec_role", t=title))
                c1, c2, c3, c4 = st.columns(4)
                c1.metric(t("m_score"), f"{row['match_score']}%")
                c2.metric(t("m_postings"), f"{row['postings']}")
                c3.metric(t("m_demand"), f"{int(row['demand_score'])}/100")
                c4.metric(t("m_trend"), status_label(trend_map_me.get(title)),
                          help=(t("trend_help", v=role_shift[title]) if title in role_shift else None))

                g1, g2 = st.columns(2)
                with g1:
                    st.write(t("have_skills", s=tag_skills(row['matched']) if row['matched'] else t("generic")))
                with g2:
                    st.write(t("gap_skills", s=tag_skills(row['missing'][:6]) if row['missing'] else t("no_gap")))

                ind_txt = industry_sdg_text(row["industries"]) if isinstance(row["industries"], list) else ""
                st.markdown(t("sdg_chain", title=title, sdg=sdg_label(sdgs), desc=sdg_desc)
                            + (t("sdg_industry", s=ind_txt) if ind_txt else ""))
                if row["sdg_hit"]:
                    st.caption(t("sdg_match", s=", ".join(sdg_name(n) for n in row["sdg_hit"])))

            # ---- 再培训路径 + 学习进度与徽章 ----
            steps, total, weeks, ph, cvh = build_plan(row["missing"], weekly_hours, skill_df, signals)
            with st.expander(t("plan_title", title=title, total=total, weeks=weeks)):
                st.write(t("plan_time", total=total, wh=weekly_hours, weeks=weeks))
                nores = no_resource_skills(row["missing"][:6])
                if nores:
                    st.warning(t("plan_nores", s=", ".join(nores)))
                if not steps:
                    st.info(t("plan_empty"))
                if skill_df is not None and steps:
                    st.caption(t("plan_prior"))
                if signals and steps:
                    parts = []
                    for sk, *_ in steps:
                        _, sh, h = skill_momentum(sk, skill_df, signals)
                        parts.append(f"{sk}（{sh:+.1f}pp ｜ {h}/{len(signals)}）")
                    st.caption(t("plan_why", s="；".join(parts)))
                done_flags, done_hours, learn_done, skill_badges = [], 0, 0, []
                for sk, name, url, cost, hrs in steps:
                    label = f"{len(done_flags) + 1}\\. " + t("plan_skill", sk=sk, name=name, url=url, cost=t(cost), h=hrs,
                                                              sdg='/'.join(map(str, skill_sdg(sk))))
                    ok = step_check(label, f"prog_{title}_{sk}")
                    done_flags.append(ok)
                    if ok:
                        done_hours += hrs
                        learn_done += 1
                        skill_badges.append(sk)
                ok = step_check(f"{len(done_flags) + 1}\\. " + t("plan_project", title=title, h=ph), f"prog_{title}_project")
                done_flags.append(ok)
                done_hours += ph if ok else 0
                ok = step_check(f"{len(done_flags) + 1}\\. " + t("plan_cv", h=cvh), f"prog_{title}_cv")
                done_flags.append(ok)
                done_hours += cvh if ok else 0

                n_done, n_all = sum(done_flags), len(done_flags)
                st.markdown(t("prog_title"))
                st.progress(n_done / n_all)
                st.caption(t("prog_text", d=n_done, n=n_all, h=max(0, total - done_hours)))
                badges = [t("badge_skill", s=s_) for s_ in skill_badges]
                if n_done >= 1:
                    badges.append(t("badge_start"))
                if steps and learn_done == len(steps):
                    badges.append(t("badge_builder"))
                if n_done == n_all:
                    badges.append(t("badge_ready"))
                st.write("  \n".join(badges) if badges else t("badge_none"))
                st.caption(t("plan_cap", p=learn_pref, g=growth_mindset))
                st.caption(t("plan_note"))
                st.markdown(t("hrd"))

            # ---- 次要信息（与上面的 expander 同级；Streamlit 不允许 expander 嵌套）----
            with st.expander(t("why_h")):
                st.caption(t("score_break", c=row['coverage'] * 100, cw=100 - BG_W,
                             b=t("yes") if row['bg'] else t("no"), bw=BG_W, tb=row['trend_bonus']))
                if pref_sdgs:
                    st.caption(t("sdg_bonus_line", b=row["sdg_bonus"]))
                st.caption(t("src_job", n=row['postings']))
                if signals:
                    if row["trend_hits"] > 0:
                        st.caption(t("sig_hit", n=len(signals), h=row['trend_hits'], ts=fetched_at))
                    else:
                        st.caption(t("sig_none", n=len(signals)))
                has_g = row["growth"] is not None and not pd.isna(row["growth"])
                if has_g:
                    base_g = row["growth"]
                    st.markdown(t("scen_title", g=base_g))
                    s0, s1, s2 = st.columns(3)
                    for col, name, m in [(s0, t("scen_slow"), mult_slow), (s1, t("scen_mod"), mult_mod), (s2, t("scen_fast"), mult_fast)]:
                        g = base_g * m
                        col.metric(name, f"{g:+.1f}% YoY", help=t("scen_help"))
                        a_, b_ = g * (1 - spread / 100), g * (1 + spread / 100)
                        col.caption(t("scen_range", lo=min(a_, b_), hi=max(a_, b_)))
                    st.caption(t("scen_warn"))
                else:
                    st.caption(t("no_growth"))

            top_sdg = ", ".join(sdg_name(n) for n, _ in chosen_sdgs.most_common(3) if n in SDG_NAMES)
            st.markdown(t("sdg_overall", s=top_sdg))

    # ======================= 市场趋势 =======================
    with sub_trend:
        render_radar(role_df, skill_df, trend_meta, user_skills, signals)
        render_live_panel(skill_df)  # [NEW]

    # ======================= 公平与来源 =======================
    with sub_trust:
        render_fairness(user_skills, jobs_df, BG_W)
        render_sdg_coverage(jobs_df, vocab)

with tab_team:
    render_team_mode(jobs_df, role_df, skill_df, signals, BG_W)